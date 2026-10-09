"""Reader for the hand-drawn "Pomocníci v místnostech" table (.xlsx): the
Buildings, their Rooms and how many Helpers each solver Role needs in each,
read off the sheet's merged cells and colours rather than from any number in it.

The sheet's shape (see ``tests/test_building_sheet.py`` for a small one):

* Two header rows above the first labelled row: the Buildings (a merged cell
  over the columns of its Rooms) and, under them, the Rooms. A Room that spans
  two columns (a large one) is a merged header cell, or a header cell followed
  by blank ones.
* Column A labels the rows. A row group is a Role (``Opravovatelé``, ``Měniči``,
  ``Skenovači``, ``Kresliči``, ``Fotografové``, ``Záloha``) when its label starts
  like the Role's name, and covers the rows its (merged) label cell covers. Any
  other label (the Organizer and Additional roles) is ignored.
* Under a Room, a coloured cell in a Role's rows is one Helper needed there and a
  gray, white or empty one is not. A merged cell is one slot however many cells it
  covers: when it sits under a single Room it counts for that Room, when it
  spans several (the Fotograf rows) it counts for the whole Building.

Beyond the counts, the sheet also carries who leads and how its cells are merged:

* The leadership rows (``Vedoucí budovy``, ``Pravá ruka``, ``Vedoucí místností``,
  ``Technická podpora``) hold Organizers' names, several in one cell separated by
  commas. Each non-empty cell becomes a :class:`SlotCell`: Vedoucí budovy and
  Technická podpora at the Building, Pravá ruka and Vedoucí místností under the
  first Room the cell covers.
* A cell merged sideways over several Rooms in a Room-level row (any Role, Pravá
  ruka, Vedoucí místností, Focení předávání cen, Uvaděči účastníků) is a sideways
  merge of those Rooms in that row (``cell_merges``).
* A cell merged top-to-bottom over two rows that may form a tall cell (see
  ``rostering.row_merges``) is a tall merge (``row_merges``); its names, written in
  its top-left cell, belong to both roles, a cell over Vedoucí budovy filed at the
  Building for both.

Apart from the names of the leadership rows, text written in the cells is never
read, nor any number.
"""
from __future__ import annotations

import colorsys
import io
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX, Color

from rostering.domain import Role, normalize_name
from rostering.row_merges import TALL_PAIRS

# A label belongs to the Role whose stem the normalized label starts with
# ("Opravovatelé" -> "opravovatel").
_ROLE_STEMS: dict[str, Role] = {
    "opravovatel": Role.Opravovatel,
    "menic": Role.Menic,
    "skenovac": Role.Skenovac,
    "kreslic": Role.Kreslic,
    "fotograf": Role.Fotograf,
    "zaloha": Role.Zaloha,
}

# Every row group the sheet may carry, by the stem its normalized label starts
# with: a Role name, a leadership slot or an Additional role (as grid row keys).
_ROW_STEMS: dict[str, str] = {
    **{stem: role.name for stem, role in _ROLE_STEMS.items()},
    "vedoucibudov": "VedouciBudovy",
    "pravaruk": "PravaRuka",
    "vedoucimistnost": "VedouciMistnosti",
    "technickapodpor": "TechnickaPodpora",
    "foceni": "FoceniPredavaniCen",
    "uvadec": "UvadeciUcastniku",
    "registrac": "Registrace",
}
# Leadership rows whose cells hold Organizers' names, in the order they are read.
_SLOT_KEYS = ("VedouciBudovy", "PravaRuka", "VedouciMistnosti", "TechnickaPodpora")
# The rows with one cell per Room (their sideways merges are read).
_ROOM_ROW_KEYS = tuple(role.name for role in Role) + (
    "PravaRuka",
    "VedouciMistnosti",
    "FoceniPredavaniCen",
    "UvadeciUcastniku",
)
# Slots held at the Building, whatever the cell covers.
_BUILDING_SLOTS = ("VedouciBudovy", "TechnickaPodpora")

# Channels of a colour closer together than this make it a gray (or white).
_GRAY_SPREAD = 12

# The default Office palette, for a workbook whose theme can't be read: indexed
# by the sheet's theme numbers (light 1, dark 1, light 2, dark 2, accents 1-6).
_OFFICE_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]
_THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@dataclass
class SlotCell:
    """The names written in one cell of a leadership row: ``role`` (a
    ``StructuralRole`` name) at ``building`` and, for a Room-level slot, ``room``
    (the first Room the cell covers); ``room`` is None for a Building-level one."""

    role: str
    building: str
    room: Optional[str]
    names: list[str]


@dataclass
class BuildingSheet:
    """The layout the sheet describes, in the saved config's shape
    (``[{name, rooms: [{name, capacities}], capacities}]``), the leadership cells,
    the sideways merges (``cell_merges``: row key -> Building -> Room pairs) and the
    tall merges (``row_merges``: ``{building, room, row}``), and what the user
    should look at: Czech lines for what could not be read."""

    buildings: list[dict]
    warnings: list[str] = field(default_factory=list)
    slots: list[SlotCell] = field(default_factory=list)
    cell_merges: dict[str, dict[str, list[list[str]]]] = field(default_factory=dict)
    row_merges: list[dict] = field(default_factory=list)


@dataclass
class _Layout:
    """One Building as the sheet draws it: its columns and its Rooms."""

    name: str
    first: int
    last: int
    rooms: list[tuple[str, int, int]]  # (name, first column, last column)

    def covered(self, first: int, last: int) -> list[str]:
        """The Rooms whose columns overlap ``first``..``last``."""
        return [name for name, a, b in self.rooms if a <= last and first <= b]


def parse_building_sheet(content: bytes) -> BuildingSheet:
    """Read the Buildings, Rooms and Role counts off the first sheet of an
    .xlsx file; a ``ValueError`` (Czech text) says why a file isn't one."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content))
    except Exception as exc:  # openpyxl raises a zoo of types for a non-xlsx
        raise ValueError("Soubor se nepodařilo přečíst jako tabulku .xlsx.") from exc
    return _Reader(workbook).read()


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


class _Reader:
    def __init__(self, workbook: openpyxl.Workbook) -> None:
        self.ws = workbook.active
        self.theme = _theme_colours(getattr(workbook, "loaded_theme", None))
        # Every cell of a merged range -> that range's (min_row, min_col, max_row, max_col).
        self.spans: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for rng in self.ws.merged_cells.ranges:
            box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.spans[(r, c)] = box

    # ------------------------------------------------------------------ cells
    def _box(self, row: int, col: int) -> tuple[int, int, int, int]:
        return self.spans.get((row, col), (row, col, row, col))

    def _text(self, row: int, col: int) -> str:
        anchor_row, anchor_col, _, _ = self._box(row, col)
        return _clean(self.ws.cell(anchor_row, anchor_col).value)

    def _needed(self, row: int, col: int) -> bool:
        """Whether the cell (its merged range's top-left one) is coloured."""
        anchor_row, anchor_col, _, _ = self._box(row, col)
        fill = self.ws.cell(anchor_row, anchor_col).fill
        if not fill or not fill.fill_type:
            return False
        rgb = self._rgb(fill.fgColor)
        return rgb is not None and max(rgb) - min(rgb) > _GRAY_SPREAD

    def _rgb(self, color: Optional[Color]) -> Optional[tuple[int, int, int]]:
        if color is None:
            return None
        if color.type == "rgb" and isinstance(color.rgb, str) and re.fullmatch(r"[0-9A-Fa-f]{6,8}", color.rgb):
            return _hex_to_rgb(color.rgb[-6:])
        if color.type == "indexed" and isinstance(color.indexed, int) and color.indexed < len(COLOR_INDEX):
            return _hex_to_rgb(COLOR_INDEX[color.indexed][-6:])
        if color.type == "theme" and isinstance(color.theme, int) and color.theme < len(self.theme):
            return _tinted(_hex_to_rgb(self.theme[color.theme]), color.tint or 0.0)
        return None

    # ------------------------------------------------------------------ layout
    def read(self) -> BuildingSheet:
        ws = self.ws
        labelled = self._labelled_rows()
        role_rows = {Role[key]: rows for key, rows in labelled.items() if key in Role.__members__}
        if not role_rows:
            raise ValueError(
                "V tabulce jsem nenašel řádky rolí (Opravovatelé, Měniči, Skenovači, Kresliči…) v prvním sloupci."
            )
        first_label = min(r for r in range(1, ws.max_row + 1) if self._text(r, 1))
        headers = [
            r for r in range(1, first_label) if any(self._text(r, c) for c in range(2, ws.max_column + 1))
        ]
        if len(headers) < 2:
            raise ValueError("Nad řádky rolí musí být dva řádky záhlaví: budovy a pod nimi místnosti.")
        building_row, room_row = headers[-2], headers[-1]

        warnings: list[str] = []
        buildings: list[dict] = []
        layouts: list[_Layout] = []
        seen: set[str] = set()
        for name, first, last in self._building_spans(building_row, room_row):
            if normalize_name(name) in seen:
                raise ValueError(f"Budova „{name}“ je v tabulce víckrát.")
            seen.add(normalize_name(name))
            rooms = self._rooms(room_row, first, last)
            if not rooms:
                warnings.append(f"Budova „{name}“ nemá v záhlaví žádné místnosti, přeskočena.")
                continue
            buildings.append(self._building(name, rooms, role_rows, warnings))
            layouts.append(_Layout(name, rooms[0][1], rooms[-1][2], rooms))
        if not buildings:
            raise ValueError("V tabulce jsem nenašel žádnou budovu s místnostmi.")
        tall = self._tall_boxes(labelled, layouts, warnings)
        return BuildingSheet(
            buildings=buildings,
            warnings=warnings,
            slots=self._slots(labelled, layouts, tall, warnings),
            cell_merges=self._cell_merges(labelled, layouts),
            row_merges=[
                {"building": info["building"].name, "room": info["rooms"][0], "row": info["keys"][0]}
                for info in tall.values()
            ],
        )

    def _labelled_rows(self) -> dict[str, list[int]]:
        """Each known row group's rows by key (a Role name, a leadership slot or an
        Additional role): those its label cell covers (all of a merged one)."""
        result: dict[str, list[int]] = {}
        for row in range(1, self.ws.max_row + 1):
            anchor_row, _, last_row, _ = self._box(row, 1)
            if anchor_row != row:
                continue
            label = normalize_name(self._text(row, 1))
            for stem, key in _ROW_STEMS.items():
                if label.startswith(stem):
                    result.setdefault(key, []).extend(range(row, last_row + 1))
                    break
        return result

    def _building_spans(self, building_row: int, room_row: int) -> list[tuple[str, int, int]]:
        """(name, first column, last column) of each Building: its merged header
        cell's columns, else up to the next Building (or the last Room column)."""
        ws = self.ws
        starts = [
            c for c in range(2, ws.max_column + 1) if self._text(building_row, c) and self._box(building_row, c)[1] == c
        ]
        last_room_col = max(
            (c for c in range(2, ws.max_column + 1) if self._text(room_row, c)), default=ws.max_column
        )
        spans = []
        for i, col in enumerate(starts):
            merged_end = self._box(building_row, col)[3]
            following = starts[i + 1] - 1 if i + 1 < len(starts) else last_room_col
            spans.append((self._text(building_row, col), col, merged_end if merged_end > col else following))
        return spans

    def _rooms(self, room_row: int, first: int, last: int) -> list[tuple[str, int, int]]:
        """(name, first column, last column) of each Room under a Building. A
        blank header cell continues the Room before it (a large Room in two
        columns); one before any Room is no Room."""
        rooms: list[list] = []
        col = first
        while col <= last:
            c2 = self._box(room_row, col)[3]
            name = self._text(room_row, col)
            end = min(c2, last)
            if name:
                rooms.append([name, col, end])
            elif rooms:
                rooms[-1][2] = end
            col = end + 1
        return [(n, a, b) for n, a, b in rooms]

    # ------------------------------------------------------------------ leadership and merges
    def _tall_boxes(
        self, labelled: dict[str, list[int]], layouts: list[_Layout], warnings: list[str]
    ) -> dict[tuple[int, int], dict]:
        """The cells merged top-to-bottom over two rows that may form a tall cell,
        by the (row, column) of their top-left cell: ``keys`` (upper, lower),
        ``building`` (its :class:`_Layout`) and ``rooms`` (the Rooms it covers).
        A cell over rows that can't be a tall cell, or one over Vedoucí budovy that
        doesn't span its whole Building, is reported and left out."""
        key_of_row = {row: key for key, rows in labelled.items() for row in rows}
        found: dict[tuple[int, int], dict] = {}
        for box in sorted(set(self.spans.values())):
            first_row, first_col, last_row, last_col = box
            if last_row == first_row or first_col < 2:
                continue
            keys = list(dict.fromkeys(key_of_row[r] for r in range(first_row, last_row + 1) if r in key_of_row))
            if len(keys) < 2:
                continue
            where = f"{openpyxl.utils.get_column_letter(first_col)}{first_row}"
            layout = next((l for l in layouts if l.first <= first_col <= l.last), None)
            if layout is None:
                continue
            rooms = layout.covered(max(first_col, layout.first), min(last_col, layout.last))
            if len(keys) != 2 or TALL_PAIRS.get(keys[0]) != keys[1]:
                warnings.append(f"Buňka {where} spojuje řádky, které se spojit nedají ({', '.join(keys)}); vynechána.")
            elif keys[0] == "VedouciBudovy" and len(rooms) != len(layout.rooms):
                warnings.append(f"Buňka {where} spojuje Vedoucí budovy s Pravou rukou, ale nepokrývá celou budovu; vynechána.")
            elif rooms:
                found[(first_row, first_col)] = {"keys": tuple(keys), "building": layout, "rooms": rooms}
        return found

    def _slots(
        self,
        labelled: dict[str, list[int]],
        layouts: list[_Layout],
        tall: dict[tuple[int, int], dict],
        warnings: list[str],
    ) -> list[SlotCell]:
        slots: list[SlotCell] = []
        for key in _SLOT_KEYS:
            for row in labelled.get(key, []):
                seen: set[tuple[int, int]] = set()
                for col in range(2, self.ws.max_column + 1):
                    first_row, first_col, _, last_col = self._box(row, col)
                    if (first_row, first_col) in seen or first_row != row:
                        continue  # the lower part of a cell whose names sit above
                    seen.add((first_row, first_col))
                    names = [n.strip() for n in re.split(r"[,;\n]+", str(self.ws.cell(row, first_col).value or "")) if n.strip()]
                    if not names:
                        continue
                    where = f"{openpyxl.utils.get_column_letter(first_col)}{row}"
                    layout = next((l for l in layouts if l.first <= first_col <= l.last), None)
                    rooms = layout.covered(max(first_col, layout.first), min(last_col, layout.last)) if layout else []
                    if layout is None or not rooms:
                        warnings.append(f"Jména v buňce {where} nepatří k žádné místnosti, přeskočena.")
                        continue
                    keys = tall[(row, first_col)]["keys"] if (row, first_col) in tall else (key,)
                    for slot in keys:
                        at_building = slot in _BUILDING_SLOTS or "VedouciBudovy" in keys
                        slots.append(SlotCell(slot, layout.name, None if at_building else rooms[0], names))
        return slots

    def _cell_merges(self, labelled: dict[str, list[int]], layouts: list[_Layout]) -> dict[str, dict[str, list[list[str]]]]:
        """The sideways merges of every Room-level row: a pair of neighbouring Rooms
        is merged in a row when a merged cell covers both in every row of that
        row group."""
        result: dict[str, dict[str, list[list[str]]]] = {}
        for key in _ROOM_ROW_KEYS:
            rows = labelled.get(key)
            if not rows:
                continue
            for layout in layouts:
                order = {name: i for i, (name, _, _) in enumerate(layout.rooms)}
                per_row = []
                for row in rows:
                    pairs: set[tuple[str, str]] = set()
                    for col in range(layout.first, layout.last + 1):
                        _, first_col, _, last_col = self._box(row, col)
                        covered = layout.covered(max(first_col, layout.first), min(last_col, layout.last))
                        pairs.update(zip(covered, covered[1:]))
                    per_row.append(pairs)
                merged = set.intersection(*per_row)
                if merged:
                    result.setdefault(key, {})[layout.name] = [list(pair) for pair in sorted(merged, key=lambda pr: order[pr[0]])]
        return result

    # ------------------------------------------------------------------ counts
    def _building(
        self, name: str, rooms: list[tuple[str, int, int]], role_rows: dict[Role, list[int]], warnings: list[str]
    ) -> dict:
        first, last = rooms[0][1], rooms[-1][2]
        room_caps: list[dict[str, dict]] = [{} for _ in rooms]
        building_caps: dict[str, dict] = {}
        for role, rows in role_rows.items():
            room_counts = [0] * len(rooms)
            building_count = 0
            seen: set[tuple[int, int]] = set()
            for row in rows:
                for col in range(first, last + 1):
                    anchor_row, anchor_col, _, end_col = self._box(row, col)
                    if (anchor_row, anchor_col) in seen:
                        continue
                    seen.add((anchor_row, anchor_col))
                    if not self._needed(row, col):
                        continue
                    covered = [i for i, (_, a, b) in enumerate(rooms) if a <= end_col and anchor_col <= b]
                    if len(covered) == 1:
                        room_counts[covered[0]] += 1
                    elif len(covered) > 1:
                        building_count += 1
            for caps, count in zip(room_caps, room_counts):
                if count:
                    caps[role.name] = {"minimum": count}
            if building_count:
                building_caps[role.name] = {"minimum": building_count}
        for (room_name, _, _), caps in zip(rooms, room_caps):
            if not caps:
                warnings.append(f"Místnost „{room_name}“ (budova {name}) nemá v tabulce žádnou barevnou buňku.")
        return {
            "name": name,
            "rooms": [{"name": n, "capacities": caps} for (n, _, _), caps in zip(rooms, room_caps)],
            "capacities": building_caps,
        }


# ---------------------------------------------------------------------- colours
def _hex_to_rgb(value: str) -> tuple[int, int, int]:
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def _tinted(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    """A theme colour lightened (tint > 0) or darkened (tint < 0) as Excel does."""
    if not tint:
        return rgb
    h, lightness, s = colorsys.rgb_to_hls(*(v / 255 for v in rgb))
    lightness = lightness * (1 + tint) if tint < 0 else lightness + (1 - lightness) * tint
    return tuple(round(v * 255) for v in colorsys.hls_to_rgb(h, lightness, s))  # type: ignore[return-value]


def _theme_colours(theme_xml: Optional[bytes]) -> list[str]:
    """The workbook theme's ten colours in the order theme numbers index them,
    the Office defaults for a file without a readable theme."""
    if not theme_xml:
        return _OFFICE_THEME
    try:
        scheme = ET.fromstring(theme_xml).find(f".//{_DRAWING_NS}clrScheme")
        colours = []
        for key in _THEME_ORDER:
            node = scheme.find(f"{_DRAWING_NS}{key}")[0]
            colours.append(node.get("val") if node.tag.endswith("srgbClr") else node.get("lastClr"))
        return colours if all(colours) else _OFFICE_THEME
    except (ET.ParseError, AttributeError, TypeError, IndexError):
        return _OFFICE_THEME
