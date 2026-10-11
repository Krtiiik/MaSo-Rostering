"""A fictional "Pomocníci v místnostech" table for the manual, and a picture of
it.

The Buildings tab reads such a table (``rostering/ingest/building_sheet.py``):
Buildings and Rooms from the two header rows, the number of Helpers a Role needs
from the coloured cells, and the leaders from the names written in the leadership
rows. ``build`` writes one with openpyxl so the screenshots script can load it;
``html`` draws the same workbook as an Excel-like page, because the manual shows
the reader what their own table should look like and Excel itself cannot be
driven here.
"""
from __future__ import annotations

from html import escape
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter

NEEDED = "FFE599"  # any colour but gray: one Helper needed
GRAY = "D9D9D9"  # gray (or white, or empty): none needed

# Building -> Rooms. The same layout the manual's other screenshots use.
BUILDINGS = {
    "Malá Strana": ["MS 1", "MS 2"],
    "Karlov": ["M 1", "M 2"],
    "Karlín": ["K 1", "K 2"],
}
# Role rows: label -> how many rows it has (one coloured cell per row and Room).
ROLE_ROWS = {
    "Opravovatelé": 2,
    "Měniči": 1,
    "Skenovači": 1,
    "Kresliči": 1,
    "Fotografové": 1,
    "Záloha": 1,
}

_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build(path: Path, building_leaders: dict[str, str], room_leaders: dict[str, str]) -> None:
    """Write the table. ``building_leaders``: Building -> Organizer's name (Vedoucí
    budovy, one merged cell over the Building); ``room_leaders``: Room -> name
    (Vedoucí místností)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pomocníci v místnostech"
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def put(cell: str, value=None, fill: str | None = None, bold: bool = False) -> None:
        ws[cell] = value
        ws[cell].alignment = center
        ws[cell].border = _BOX
        if fill:
            ws[cell].fill = PatternFill("solid", fgColor=fill)
        if bold:
            ws[cell].font = openpyxl.styles.Font(bold=True)

    column = 2
    spans: list[tuple[str, int, list[tuple[str, int]]]] = []
    for building, rooms in BUILDINGS.items():
        spans.append((building, column, [(room, column + i) for i, room in enumerate(rooms)]))
        column += len(rooms)
    last_column = column - 1

    # Row 1 Buildings (merged over their Rooms), row 2 Rooms.
    for building, first, rooms in spans:
        last = first + len(rooms) - 1
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
        put(f"{get_column_letter(first)}1", building, bold=True)
        for c in range(first, last + 1):
            ws.cell(1, c).border = _BOX
        for room, c in rooms:
            put(f"{get_column_letter(c)}2", room, bold=True)

    # Row 3 Vedoucí budovy, row 4 Vedoucí místností: names of Organizers.
    put("A3", "Vedoucí budovy", bold=True)
    put("A4", "Vedoucí místností", bold=True)
    for building, first, rooms in spans:
        last = first + len(rooms) - 1
        ws.merge_cells(start_row=3, start_column=first, end_row=3, end_column=last)
        put(f"{get_column_letter(first)}3", building_leaders.get(building))
        for c in range(first, last + 1):
            ws.cell(3, c).border = _BOX
        for room, c in rooms:
            put(f"{get_column_letter(c)}4", room_leaders.get(room))

    # Then the Roles: a coloured cell is one Helper needed, a gray one is none.
    row = 5
    for label, height in ROLE_ROWS.items():
        if height > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row + height - 1, end_column=1)
        put(f"A{row}", label, bold=True)
        for r in range(row, row + height):
            ws.cell(r, 1).border = _BOX
            for c in range(2, last_column + 1):
                put(f"{get_column_letter(c)}{r}", fill=GRAY if label == "Záloha" else NEEDED)
        row += height

    ws.column_dimensions["A"].width = 20
    for c in range(2, last_column + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    wb.save(path)


def html(path: Path) -> str:
    """The workbook as an Excel-like page: column letters, row numbers, merged
    cells, fills. Cells carry ``data-cell`` (``B1``) so the screenshots script can
    point at them."""
    ws = openpyxl.load_workbook(path).active
    skipped: set[tuple[int, int]] = set()
    spans: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    skipped.add((r, c))

    head = "".join(f"<th>{get_column_letter(c)}</th>" for c in range(1, ws.max_column + 1))
    rows = []
    for r in range(1, ws.max_row + 1):
        cells = [f"<th>{r}</th>"]
        for c in range(1, ws.max_column + 1):
            if (r, c) in skipped:
                continue
            cell = ws.cell(r, c)
            rowspan, colspan = spans.get((r, c), (1, 1))
            fill = cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type else "FFFFFF"
            weight = "bold" if cell.font and cell.font.bold else "normal"
            cells.append(
                f'<td data-cell="{get_column_letter(c)}{r}" rowspan="{rowspan}" colspan="{colspan}" '
                f'style="background:#{fill};font-weight:{weight}">{escape(str(cell.value or ""))}</td>'
            )
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return f"""<!doctype html><meta charset="utf-8">
<style>
  body {{ margin: 0; padding: 24px; background: #fff; font: 14px Calibri, 'Segoe UI', sans-serif; }}
  table {{ border-collapse: collapse; }}
  th {{ background: #f2f2f2; color: #555; font-weight: normal; border: 1px solid #bfbfbf; padding: 3px 10px; }}
  td {{ border: 1px solid #bfbfbf; padding: 8px 10px; text-align: center; min-width: 96px; height: 24px; }}
  td:first-of-type {{ min-width: 150px; }}
</style>
<table><tr><th></th>{head}</tr>{"".join(rows)}</table>"""
