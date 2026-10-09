"""PROTOTYPE (throwaway) — issue #16 "Replay past Seasons with old vs new scoring".

QUESTION
    Does the candidate role-Preference scoring from issue #15 actually produce
    better rosters than today's? Replay every past Season under the current
    formula and under the candidate, on identical inputs, and compare:
      * distribution of assigned Preference levels
      * single-"Ano" Helpers who got their Role
      * multi-"Ano" Helpers who got one of theirs
      * what blank answers got
      * Záloha usage
      * side effects on Building and friend satisfaction

RUN (from the repo root; needs the project venv + the gitignored data/ dir)
    python prototypes/replay-scoring/replay_scoring_PROTOTYPE.py
    python prototypes/replay-scoring/replay_scoring_PROTOTYPE.py --quick   # 1 season, 3 variants, 10 s
  -> prototypes/replay-scoring/out/results.json + out/results.html (aggregates only, no names)

ASSUMPTIONS (the replay's own, not the ticket's)
    * Only 2026-jaro has a real config.yaml. For every other Season the room
      layout and per-role minimums are derived from that Season's hand-built
      roster (data/rosters/*.xlsx), counting names per room and role. This is
      exact for 2026-jaro (the real config equals the roster counts), so it
      is assumed a fair stand-in for the older configs. If a derived config
      needs more Helpers than the Season had, every minimum is scaled down
      uniformly (reported in the output).
    * The solver here is a FORK of rostering/solver/model.py that takes the
      role cost as a function. Variant "old" must reproduce production's
      objective; the script asserts that on 2026-jaro (see --no-fidelity).
    * All objective terms are multiplied by 60 (= lcm(1..5)) so the candidate's
      (1 + 1/k) factor stays an exact integer for k = 1..5.
    * Building-mismatch (3) and friend (5) weights stay at production values;
      the candidate's unit ``u`` (cost per relative unit) is swept to see how
      the role penalty's absolute scale trades off against them. u = 4 keeps
      today's cost for a Nevadí at k = 0 (2 * 4 = 8 = old (5-3) * 4).

The pure scoring functions are in the "CANDIDATE SCORING" section and have no
solver / DOM dependencies, so they can be lifted into rostering/solver/.
"""
from __future__ import annotations

import argparse
import glob
import html
import json
import math
import os
import subprocess
import sys
import time
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402
from ortools.sat.python import cp_model  # noqa: E402

from rostering.config import load_buildings  # noqa: E402
from rostering.domain import (  # noqa: E402
    Building, Competition, Helper, Preference, Role, RoleCapacity, Room, normalize_name,
)
from rostering.ingest.mapping import building_keys  # noqa: E402
from rostering.ingest.raw_survey import parse_raw_survey  # noqa: E402
from rostering.solver.model import SolverConfig, SolverWeights, solve_competition  # noqa: E402
from rostering.solver.scoring import FriendScoringConfig, build_friend_pairs  # noqa: E402

SEASONS = ["2023-podzim", "2024-jaro", "2024-podzim", "2025-jaro", "2025-podzim", "2026-jaro"]
SCALE = 60  # lcm(1..5): makes cost * (k + 1) / k an exact integer for k = 1..5
OUT = Path(__file__).resolve().parent / "out"

# ════════════════════════════════════════════════════════════════════════════
# CANDIDATE SCORING  (pure — the part worth lifting into the real codebase)
# ════════════════════════════════════════════════════════════════════════════

CANDIDATE_TABLE = {  # issue #15: relative cost per rating
    Preference.Ano: 0,
    Preference.Klidne: 1,
    Preference.Nevadi: 2,
    Preference.Spise_ne: 6,
    Preference.Ne: 12,
}
CANDIDATE_ZALOHA = 4


@dataclass(frozen=True)
class Scoring:
    key: str
    label: str
    kind: str  # "old" | "new"
    unit: int = 4  # integer cost per relative unit (new only)
    zaloha: int = CANDIDATE_ZALOHA  # relative cost of Záloha (new only)
    blank_is_nevadi: bool = True  # False: a blank costs 0, as today (new only)
    concentration: bool = True  # False: skip the (1 + 1/k) factor (new only)


def count_ano(h: Helper) -> int:
    return sum(1 for p in h.role_preferences.values() if p == Preference.Ano)


def role_costs(h: Helper, s: Scoring, old_role_weight: int = 4) -> dict[Role, int]:
    """Integer cost (× SCALE) of assigning ``h`` each of the 6 Roles."""
    if s.kind == "old":
        # Today: (5 - pref) * weight for a rated role; blank and Záloha cost 0.
        max_pref = max(p.value for p in Preference)
        return {
            r: (max_pref - int(h.role_preferences[r])) * old_role_weight * SCALE
            if r in h.role_preferences else 0
            for r in Role
        }
    k = count_ano(h)
    factor = SCALE * (k + 1) // k if (s.concentration and k > 0) else SCALE  # (1 + 1/k) * SCALE
    out: dict[Role, int] = {}
    for r in Role:
        if r == Role.Zaloha:
            rel = s.zaloha
        elif r in h.role_preferences:
            rel = CANDIDATE_TABLE[h.role_preferences[r]]
        else:
            rel = CANDIDATE_TABLE[Preference.Nevadi] if s.blank_is_nevadi else 0
        out[r] = rel * s.unit * factor if rel else 0
    return out


def make_variants() -> list[Scoring]:
    return [
        Scoring("old", "today (blank=0, Záloha=0)", "old"),
        Scoring("new_u1", "candidate, u=1", "new", unit=1),
        Scoring("new_u2", "candidate, u=2", "new", unit=2),
        Scoring("new_u4", "candidate, u=4 (anchor)", "new", unit=4),
        Scoring("new_u8", "candidate, u=8", "new", unit=8),
        Scoring("abl_noconc", "u=4 without concentration factor", "new", unit=4, concentration=False),
        Scoring("abl_nozal", "u=4 with free Záloha", "new", unit=4, zaloha=0),
        Scoring("abl_blank0", "u=4 with blank=0 (as today)", "new", unit=4, blank_is_nevadi=False),
    ]


# ════════════════════════════════════════════════════════════════════════════
# SEASON INPUTS
# ════════════════════════════════════════════════════════════════════════════

def find_data_dir(arg: str | None) -> Path:
    if arg:
        return Path(arg)
    if (REPO / "data").is_dir():
        return REPO / "data"
    common = subprocess.check_output(
        ["git", "rev-parse", "--git-common-dir"], cwd=REPO, text=True).strip()
    main_root = (REPO / common).resolve().parent
    if (main_root / "data").is_dir():
        return main_root / "data"
    sys.exit("data/ directory not found; pass --data-dir")


ROLE_ROWS = {
    "opravovatele": Role.Opravovatel, "menici": Role.Menic,
    "skenovaci": Role.Skenovac, "kreslici": Role.Kreslic,
}

ROSTER_FILES = {  # substring of the hand-built roster's file name per Season
    "2023-podzim": "Počty pomocníků Praha podzim 2023",
    "2024-jaro": "Rozřazení pomocníků jaro 2024",
    "2024-podzim": "Rozdělení pomocníků Praha - podzim 2024",
    "2025-jaro": "Rozdělení pomocníků Praha - jaro 2025",
    "2025-podzim": "Rozdělení pomocníků Praha - podzim 2025",
    "2026-jaro": "Rozdělení pomocníků Praha - jaro 2026",
}


def derive_buildings(roster_path: str) -> dict[str, Building]:
    """Count names per (room, role) in a hand-built roster sheet."""
    wb = openpyxl.load_workbook(roster_path)
    sheet = "edits" if "edits" in wb.sheetnames else "Pomocníci v místnostech"
    ws = wb[sheet]
    b_cols = {c.column: str(c.value).strip() for c in ws[1] if c.value and str(c.value).strip()}
    r_cols = {c.column: str(c.value).strip() for c in ws[2] if c.value and str(c.value).strip()}
    b_cols.pop(1, None)
    r_cols.pop(1, None)
    last_col = max(r_cols) + 1  # a room spans at most 2 columns (overflow)

    def span_owner(col: int, owners: dict[int, str], end: int) -> str | None:
        keys = sorted(owners)
        hit = None
        for k in keys:
            if k <= col:
                hit = k
        if hit is None:
            return None
        nxt = next((k for k in keys if k > hit), end + 1)
        return owners[hit] if col < nxt else None

    counts: dict[tuple[str, str], int] = defaultdict(int)  # (room, role-row)
    photo: dict[str, int] = defaultdict(int)  # building
    label = None
    for row in ws.iter_rows(min_row=3, max_row=40):
        a = row[0].value
        if a and str(a).strip():
            label = normalize_name(str(a))
        for c in row[1:]:
            if c.column > last_col or not (c.value and str(c.value).strip()):
                continue
            room = span_owner(c.column, r_cols, last_col)
            bld = span_owner(c.column, b_cols, last_col)
            if label in ROLE_ROWS and room:
                counts[room, label] += 1
            elif label == "fotografove" and bld:
                photo[bld] += 1

    buildings: dict[str, Building] = {}
    b_of_room = {}
    for col, room in r_cols.items():
        b_of_room[room] = next(b_cols[k] for k in sorted(b_cols, reverse=True) if k <= col)
    for room, bname in b_of_room.items():
        b = buildings.setdefault(bname, Building(name=bname))
        b.rooms.append(Room(name=room, capacities={
            ROLE_ROWS[lab]: RoleCapacity(counts[room, lab]) for lab in ROLE_ROWS if counts[room, lab]
        }))
    for bname, n in photo.items():
        if bname in buildings:
            buildings[bname].capacities[Role.Fotograf] = RoleCapacity(n)
    return buildings


def total_minimum(buildings: dict[str, Building]) -> int:
    return sum(c.minimum for b in buildings.values()
               for c in [*b.capacities.values(), *(c for r in b.rooms for c in r.capacities.values())])


def scale_minimums(buildings: dict[str, Building], target: int) -> None:
    """Rescale the room-level minimums (largest-remainder rounding) so that all
    minimums, including the fixed building-level Fotograf ones, total ``target``."""
    caps = [c for b in buildings.values() for r in b.rooms for c in r.capacities.values() if c.minimum]
    fixed = sum(c.minimum for b in buildings.values() for c in b.capacities.values())
    base = sum(c.minimum for c in caps)
    f = max(target - fixed, 0) / base
    exact = [c.minimum * f for c in caps]
    floors = [int(x) for x in exact]
    left = (target - fixed) - sum(floors)
    for i in sorted(range(len(caps)), key=lambda i: exact[i] - floors[i], reverse=True)[:max(left, 0)]:
        floors[i] += 1
    for c, m in zip(caps, floors):
        c.minimum = m


def load_season(data_dir: Path, season: str, fill: float = 0.0) -> tuple[Competition, dict]:
    survey = parse_raw_survey(data_dir / "seasons" / season / "raw-response.xlsx")
    real_cfg = data_dir / "seasons" / season / "config.yaml"
    info: dict = {"season": season, "n_helpers": len(survey.helpers)}
    if real_cfg.exists():
        buildings = load_buildings(real_cfg)
        info["config_source"] = "config.yaml"
    else:
        path = next(p for p in glob.glob(str(data_dir / "rosters" / "*.xlsx"))
                    if normalize_name(ROSTER_FILES[season]) in normalize_name(os.path.basename(p)))
        buildings = derive_buildings(path)
        info["config_source"] = "derived from hand-built roster"
    cameras = sum(h.can_bring_camera for h in survey.helpers)
    photo_min = sum(b.capacities[Role.Fotograf].minimum for b in buildings.values() if Role.Fotograf in b.capacities)
    need = total_minimum(buildings)
    info.update(minimum_total=need, cameras=cameras, photo_min=photo_min, scaled_by=1.0)
    # Leave room for at least a few Záloha so the comparison isn't degenerate.
    limit = int(len(survey.helpers) * 0.97) - 0  # 3% slack
    if fill:  # stress regime: scale minimums so they add up to ~fill * helpers
        limit = int(len(survey.helpers) * fill)
        info["stress_fill"] = fill
    if need > limit or (fill and need < limit):
        scale_minimums(buildings, limit)
        info["scaled_by"] = round(limit / need, 3)
        info["minimum_total"] = total_minimum(buildings)
    info["buildings"] = {b.name: [r.name for r in b.rooms] for b in buildings.values()}
    info["k_ano_hist"] = dict(sorted(Counter(count_ano(h) for h in survey.helpers).items()))
    info["blank_hist"] = dict(sorted(Counter(5 - len(h.role_preferences) for h in survey.helpers).items()))
    return Competition(buildings=buildings, helpers=survey.helpers), info


# ════════════════════════════════════════════════════════════════════════════
# FORKED SOLVER — identical to rostering/solver/model.py except the role cost
# comes from ``role_costs`` and every term is multiplied by SCALE.
# ════════════════════════════════════════════════════════════════════════════

def solve(comp: Competition, s: Scoring, time_limit: float, workers: int, seed: int = 1,
          weights: SolverWeights | None = None):
    weights = weights or SolverWeights()
    friend_cfg = FriendScoringConfig()
    model = cp_model.CpModel()
    helpers, buildings, roles = comp.helpers, list(comp.buildings.values()), list(Role)

    rooms: list[tuple[str, Room]] = []
    building_rooms: dict[str, list[int]] = {}
    for b in buildings:
        building_rooms[b.name] = []
        for room in b.rooms:
            building_rooms[b.name].append(len(rooms))
            rooms.append((b.name, room))
    n_rooms = len(rooms)

    assign_room = {(h.id, r): model.NewBoolVar(f"a_{h.id}_{r}") for h in helpers for r in range(n_rooms)}
    assign_role = {(h.id, r): model.NewBoolVar(f"o_{h.id}_{r.name}") for h in helpers for r in roles}
    for h in helpers:
        model.Add(sum(assign_room[h.id, r] for r in range(n_rooms)) == 1)
        model.Add(sum(assign_role[h.id, r] for r in roles) == 1)
        if not h.can_bring_camera:
            model.Add(assign_role[h.id, Role.Fotograf] == 0)

    # role∧room linearisation, only for (role, room) pairs a minimum refers to
    needed_roles = {role for _n, rm in rooms for role, c in rm.capacities.items() if c.minimum}
    needed_roles |= {role for b in buildings for role, c in b.capacities.items() if c.minimum}
    rr = {}
    for h in helpers:
        for role in needed_roles:
            for r in range(n_rooms):
                p = model.NewBoolVar(f"rr_{h.id}_{role.name}_{r}")
                rr[h.id, role, r] = p
                model.Add(p <= assign_role[h.id, role])
                model.Add(p <= assign_room[h.id, r])
                model.Add(p >= assign_role[h.id, role] + assign_room[h.id, r] - 1)
    for rid, (_bn, room) in enumerate(rooms):
        for role, cap in room.capacities.items():
            if cap.minimum:
                model.Add(sum(rr[h.id, role, rid] for h in helpers) >= cap.minimum)
    for b in buildings:
        for role, cap in b.capacities.items():
            if cap.minimum:
                model.Add(sum(rr[h.id, role, rid] for rid in building_rooms[b.name] for h in helpers) >= cap.minimum)

    terms = []
    keys = [building_keys(bn) for bn, _r in rooms]
    for h in helpers:
        for role, cost in role_costs(h, s, weights.role_preference).items():
            if cost:
                terms.append(cost * assign_role[h.id, role])
        if h.building_preferences:
            pref = frozenset().union(*(building_keys(p) for p in h.building_preferences))
            for rid, ks in enumerate(keys):
                if ks.isdisjoint(pref):
                    terms.append(weights.building_mismatch * SCALE * assign_room[h.id, rid])
    for a, b, w in build_friend_pairs(helpers, friend_cfg):
        colo = []
        for rid in range(n_rooms):
            z = model.NewBoolVar(f"z_{a}_{b}_{rid}")
            model.Add(z <= assign_room[a, rid]); model.Add(z <= assign_room[b, rid])
            model.Add(z >= assign_room[a, rid] + assign_room[b, rid] - 1)
            colo.append(z)
        sat = model.NewBoolVar(f"fs_{a}_{b}")
        model.Add(sat == sum(colo))
        terms.append(w * weights.friend_unsatisfied * SCALE * (1 - sat))
    model.Minimize(sum(terms))

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = time_limit
    solver.parameters.num_workers = workers
    solver.parameters.random_seed = seed
    t0 = time.time()
    status = solver.Solve(model)
    took = time.time() - t0
    if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        return None
    result = {}
    for h in helpers:
        rid = next(r for r in range(n_rooms) if solver.Value(assign_room[h.id, r]))
        role = next(r for r in roles if solver.Value(assign_role[h.id, r]))
        result[h.id] = (rooms[rid][0], rooms[rid][1].name, role)
    obj, bound = solver.ObjectiveValue(), solver.BestObjectiveBound()
    return {
        "assign": result, "status": "OPTIMAL" if status == cp_model.OPTIMAL else "FEASIBLE",
        "objective": obj, "bound": bound, "gap_pct": 0.0 if obj == 0 else round(100 * (obj - bound) / obj, 2),
        "seconds": round(took, 1),
    }


# ════════════════════════════════════════════════════════════════════════════
# METRICS
# ════════════════════════════════════════════════════════════════════════════

BUCKETS = ["Ano", "Klidně", "Nevadí", "Spíš ne", "Ne", "blank", "Záloha"]
LEVEL_BUCKET = {Preference.Ano: "Ano", Preference.Klidne: "Klidně", Preference.Nevadi: "Nevadí",
                Preference.Spise_ne: "Spíš ne", Preference.Ne: "Ne"}


def evaluate(comp: Competition, res: dict) -> dict:
    assign = res["assign"]
    hs = {h.id: h for h in comp.helpers}
    levels = Counter()
    single = [0, 0]; multi = [0, 0]; multi3 = [0, 0]
    all_blank = Counter(); n_all_blank = 0
    partial = Counter(); n_partial = 0
    zal = Counter()
    for hid, (_b, _r, role) in assign.items():
        h = hs[hid]
        prefs = h.role_preferences
        k = count_ano(h)
        if role == Role.Zaloha:
            bucket = "Záloha"
        elif role in prefs:
            bucket = LEVEL_BUCKET[prefs[role]]
        else:
            bucket = "blank"
        levels[bucket] += 1
        anos = {r for r, p in prefs.items() if p == Preference.Ano}
        got_ano = role in anos
        if k == 1:
            single[1] += 1; single[0] += got_ano
        elif k >= 2:
            multi[1] += 1; multi[0] += got_ano
            if k >= 3:
                multi3[1] += 1; multi3[0] += got_ano
        n_blank = 5 - len(prefs)
        if n_blank == 5:
            n_all_blank += 1
            all_blank[role.value] += 1
        elif n_blank > 0:
            n_partial += 1
            partial["got Ano" if got_ano else "on a blank role" if bucket == "blank"
                    else "Záloha" if bucket == "Záloha" else "rated, not Ano"] += 1
        if role == Role.Zaloha:
            zal["total"] += 1
            zal["had an Ano"] += k > 0
            zal["all blank"] += n_blank == 5
            zal["Klidně/Nevadí-only (no Ano)"] += (k == 0 and n_blank < 5)

    # Building satisfaction
    bld = [0, 0]
    for hid, (bname, _r, _role) in assign.items():
        h = hs[hid]
        if h.building_preferences:
            bld[1] += 1
            pref = frozenset().union(*(building_keys(p) for p in h.building_preferences))
            bld[0] += not building_keys(bname).isdisjoint(pref)
    # Friend satisfaction (same pairwise/symmetric pairs the solver scores)
    pairs = build_friend_pairs(comp.helpers, FriendScoringConfig())
    sat = sum(1 for a, b, _w in pairs if assign[a][:2] == assign[b][:2])
    asked = {x for a, b, _w in pairs for x in (a, b)}
    happy = {x for a, b, _w in pairs if assign[a][:2] == assign[b][:2] for x in (a, b)}

    # Yardsticks: each roster's total role cost under both formulas (÷ SCALE)
    old_s, new_s = make_variants()[0], make_variants()[3]
    yard_old = sum(role_costs(hs[i], old_s)[a[2]] for i, a in assign.items()) / SCALE
    yard_new = sum(role_costs(hs[i], new_s)[a[2]] for i, a in assign.items()) / SCALE

    return {
        "levels": {b: levels[b] for b in BUCKETS},
        "single_ano": single, "multi_ano": multi, "multi_ano_k3plus": multi3,
        "all_blank": {"n": n_all_blank, "roles": dict(all_blank)},
        "partial_blank": {"n": n_partial, **partial},
        "zaloha": dict(zal),
        "building": bld, "friend_pairs": [sat, len(pairs)], "friend_helpers": [len(happy), len(asked)],
        "yard_old": round(yard_old, 1), "yard_new_u4": round(yard_new, 1),
        "status": res["status"], "gap_pct": res["gap_pct"], "seconds": res["seconds"],
    }


# ════════════════════════════════════════════════════════════════════════════
# DRIVER
# ════════════════════════════════════════════════════════════════════════════

def _job(args):
    data_dir, season, key, time_limit, workers, seed, fill = args
    comp, _info = load_season(Path(data_dir), season, fill)
    s = next(v for v in make_variants() if v.key == key)
    res = solve(comp, s, time_limit, workers, seed)
    if res is None:
        return season, key, {"status": "INFEASIBLE"}
    return season, key, evaluate(comp, res)


def fidelity_check(data_dir: Path, workers: int) -> str:
    """Variant 'old' (forked solver) must match production's objective."""
    comp, _ = load_season(data_dir, "2026-jaro")
    fork = solve(comp, make_variants()[0], 60, workers)
    prod = solve_competition(comp, SolverConfig(time_limit_seconds=60))
    if fork is None or prod is None:
        return "fidelity: one of the solves was infeasible"
    # Neither solve is proven optimal in 60 s, so allow solver noise (~5 %).
    ok = abs(fork["objective"] / SCALE - prod.objective_value) <= max(3, 0.05 * prod.objective_value)
    return (f"fidelity 2026-jaro: fork={fork['objective'] / SCALE:.1f} ({fork['status']}) vs "
            f"production={prod.objective_value:.1f} ({prod.status}) -> {'OK' if ok else 'MISMATCH'}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir")
    ap.add_argument("--seasons", nargs="*", default=SEASONS)
    ap.add_argument("--variants", nargs="*")
    ap.add_argument("--time-limit", type=float, default=60)
    ap.add_argument("--procs", type=int, default=1, help="parallel solves (CP-SAT needs the cores; keep 1)")
    ap.add_argument("--workers", type=int, default=16, help="CP-SAT workers per solve")
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--fill", type=float, default=0.0, help="stress: scale minimums to fill*helpers (e.g. 1.0)")
    ap.add_argument("--out-name", default="results", help="basename for out/<name>.json/.html")
    ap.add_argument("--quick", action="store_true")
    ap.add_argument("--no-fidelity", action="store_true")
    ap.add_argument("--configs-only", action="store_true", help="print derived configs and stop")
    args = ap.parse_args()
    data_dir = find_data_dir(args.data_dir)
    if args.quick:
        args.seasons, args.variants, args.time_limit = ["2026-jaro"], ["old", "new_u4", "abl_nozal"], 10
    keys = args.variants or [v.key for v in make_variants()]

    infos = {}
    for season in args.seasons:
        _c, info = load_season(data_dir, season, args.fill)
        infos[season] = info
        print(f"{season}: {info['n_helpers']} helpers, config {info['config_source']}, "
              f"minimums {info['minimum_total']} (scaled x{info['scaled_by']}), cameras {info['cameras']} "
              f"vs Fotograf min {info['photo_min']}")
        print(f"   buildings {info['buildings']}")
    if args.configs_only:
        return

    if not args.no_fidelity:
        print(fidelity_check(data_dir, args.workers))

    jobs = [(str(data_dir), s, k, args.time_limit, args.workers, args.seed, args.fill) for s in args.seasons for k in keys]
    results: dict[str, dict[str, dict]] = {s: {} for s in args.seasons}
    t0 = time.time()
    with ProcessPoolExecutor(max_workers=args.procs) as ex:
        for done, (season, key, m) in enumerate(ex.map(_job, jobs), 1):
            results[season][key] = m
            print(f"[{done}/{len(jobs)} {time.time() - t0:5.0f}s] {season} {key}: {m.get('status')} "
                  f"gap {m.get('gap_pct')}%", flush=True)

    OUT.mkdir(exist_ok=True)
    payload = {"variants": [v.__dict__ | {"kind": v.kind} for v in make_variants() if v.key in keys],
               "seasons": infos, "results": results, "time_limit": args.time_limit,
               "scale": SCALE}
    payload["seed"] = args.seed
    payload["fill"] = args.fill
    (OUT / f"{args.out_name}.json").write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / f"{args.out_name}.html").write_text(render_report(payload), encoding="utf-8")
    print(f"wrote {OUT / args.out_name}.json and .html")


# ════════════════════════════════════════════════════════════════════════════
# REPORT (aggregates only)
# ════════════════════════════════════════════════════════════════════════════

def _pct(n: float, d: float) -> str:
    return "–" if not d else f"{100 * n / d:.0f} %"


def render_report(p: dict) -> str:
    seasons = list(p["results"])
    keys = [v["key"] for v in p["variants"]]
    label = {v["key"]: v["label"] for v in p["variants"]}
    R = p["results"]
    ok = lambda s, k: R[s].get(k, {}).get("status") in ("OPTIMAL", "FEASIBLE")  # noqa: E731

    def pooled(k, getter):
        return [sum(getter(R[s][k])[i] for s in seasons if ok(s, k)) for i in (0, 1)]

    def ratio_table(title, note, getter):
        rows = []
        for k in keys:
            tot = pooled(k, getter)
            cells = "".join(
                f"<td>{_pct(*getter(R[s][k]))}<small>{getter(R[s][k])[0]}/{getter(R[s][k])[1]}</small></td>"
                if ok(s, k) else "<td>–</td>" for s in seasons)
            rows.append(f"<tr><th>{html.escape(label[k])}</th><td class=pool>{_pct(*tot)}<small>{tot[0]}/{tot[1]}</small></td>{cells}</tr>")
        head = "".join(f"<th>{s}</th>" for s in seasons)
        return (f"<h3>{title}</h3><p class=note>{note}</p><table><tr><th></th><th>all</th>{head}</tr>"
                + "".join(rows) + "</table>")

    def count_table(title, note, cols, getter):
        rows = []
        for k in keys:
            tot = {c: sum(getter(R[s][k]).get(c, 0) for s in seasons if ok(s, k)) for c in cols}
            grand = sum(tot.values()) or 1
            cells = "".join(f"<td>{tot[c]}<small>{100 * tot[c] / grand:.0f} %</small></td>" for c in cols)
            rows.append(f"<tr><th>{html.escape(label[k])}</th>{cells}</tr>")
        head = "".join(f"<th>{c}</th>" for c in cols)
        return f"<h3>{title}</h3><p class=note>{note}</p><table><tr><th></th>{head}</tr>" + "".join(rows) + "</table>"

    def yard_table():
        rows = []
        for k in keys:
            def avg(f):
                vals = [R[s][k][f] / p["seasons"][s]["n_helpers"] for s in seasons if ok(s, k)]
                return sum(vals) / len(vals) if vals else float("nan")
            gaps = [R[s][k]["gap_pct"] for s in seasons if ok(s, k)]
            rows.append(f"<tr><th>{html.escape(label[k])}</th><td>{avg('yard_old'):.2f}</td><td>{avg('yard_new_u4'):.2f}</td>"
                        f"<td>{max(gaps):.1f} %</td></tr>")
        return ("<h3>Role cost per Helper, read through each formula</h3><p class=note>How each roster looks "
                "through today’s lens and through the candidate (u=4) lens — lower is better in that lens. "
                "Last column: worst optimality gap among the solves (0 % = proven optimal).</p>"
                "<table><tr><th></th><th>today’s formula</th><th>candidate u=4</th><th>worst gap</th></tr>"
                + "".join(rows) + "</table>")

    blank_roles = [r.value for r in Role]
    fill = p.get("fill") or 0
    regime = (f"STRESS — minimums scaled up to {fill:.0%} of the Helpers, so Záloha has no slack and Roles compete" if fill
              else "BASELINE — minimums as configured (2026-jaro) or as in the hand-built roster; Záloha has slack")
    seasons_tbl = "".join(
        f"<tr><th>{s}</th><td>{i['n_helpers']}</td><td>{i['config_source']}</td><td>{i['minimum_total']}</td>"
        f"<td>{i['scaled_by']}</td><td>{i['k_ano_hist']}</td><td>{i['blank_hist']}</td></tr>"
        for s, i in p["seasons"].items())

    return f"""<!doctype html><html lang="cs"><head><meta charset="utf-8">
<title>Replay: old vs new role scoring — PROTOTYPE</title><style>
body{{font:15px/1.5 system-ui,sans-serif;max-width:1180px;margin:2rem auto;padding:0 1rem;color:#1c1c1c}}
h1{{margin-bottom:.2rem}}h3{{margin:2rem 0 .2rem}}.note{{color:#555;margin:.1rem 0 .5rem;max-width:80ch}}
table{{border-collapse:collapse;font-variant-numeric:tabular-nums}}td,th{{padding:.25rem .6rem;border-bottom:1px solid #ddd;text-align:right}}
th:first-child{{text-align:left;font-weight:600;white-space:nowrap}}small{{display:block;color:#777;font-size:.75em}}
.pool{{background:#eef3ff;font-weight:600}}.banner{{background:#fff4d6;padding:.5rem .8rem;border-left:4px solid #d19a00}}
</style></head><body>
<h1>Replay past Seasons: today’s vs candidate role scoring</h1>
<p class=banner>PROTOTYPE — throwaway. Issue #16. Aggregates only; solver time limit {p['time_limit']} s per solve.<br>
Regime: <b>{regime}</b></p>
<p>Every Season is re-solved on identical inputs under each scoring. “all” pools the Helpers of the Seasons shown.
CP-SAT breaks ties arbitrarily, so a difference of a couple of Helpers between two variants is noise.</p>
<h3>Inputs</h3><table><tr><th>Season</th><th>Helpers</th><th>Config</th><th>Σ minimums</th><th>scaled ×</th>
<th>#Ano histogram</th><th>#blank histogram</th></tr>{seasons_tbl}</table>
{ratio_table("Single-“Ano” Helpers who got their Role", "Helpers with exactly one Ano; share placed in that Role.", lambda m: m["single_ano"])}
{ratio_table("Multi-“Ano” Helpers who got one of theirs", "Helpers with two or more Anos; share placed in any of them.", lambda m: m["multi_ano"])}
{ratio_table("…of which three or more Anos", "Subset of the above with k ≥ 3.", lambda m: m["multi_ano_k3plus"])}
{count_table("Assigned Preference level (all Helpers)", "Preference the Helper gave to the Role they received; ‘blank’ = no answer for that Role.", BUCKETS, lambda m: m["levels"])}
{count_table("What all-blank Helpers got", "Helpers who left all five Roles blank: the Role they were assigned.", blank_roles, lambda m: m["all_blank"]["roles"])}
{count_table("What partly-blank Helpers got", "Helpers with 1–4 blanks: what they were assigned.", ["got Ano", "rated, not Ano", "on a blank role", "Záloha"], lambda m: m["partial_blank"])}
{count_table("Who ends up in Záloha", "Helpers placed in Záloha, by what they had said.", ["total", "had an Ano", "all blank", "Klidně/Nevadí-only (no Ano)"], lambda m: m["zaloha"])}
{ratio_table("Building satisfaction", "Helpers with a Building preference placed inside their acceptable set.", lambda m: m["building"])}
{ratio_table("Friend pairs satisfied", "Pairwise/symmetric pairs sharing a Room (the app’s default).", lambda m: m["friend_pairs"])}
{ratio_table("Helpers with a friend request who got at least one", "", lambda m: m["friend_helpers"])}
{yard_table()}
</body></html>"""


if __name__ == "__main__":
    main()
