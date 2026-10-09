"""Organizers as a tracked person category (see CONTEXT.md "Organizer"):
created by hand, placed only by assigning them into an Organizer role slot,
recognized across Seasons like any Person. Mutation-layer tests run against a
temp-dir workspace seeded with synthetic Helpers, Organizers and Seasons (never
anything from data/); the solver is exercised on plain domain objects, and the
export is checked by re-reading the workbook."""
import importlib
import io
from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from rostering.domain import Building, Competition, Helper, Organizer, Role, RoleCapacity, Room
from rostering.persistence.serialize import organizer_from_dict, organizer_to_dict
from rostering.persistence.workspace import Workspace
from rostering.solver.model import SolverConfig, solve_competition
from rostering.webapp import mutations

CONFIG = [
    {
        "name": "B",
        "rooms": [
            {"name": "R1", "capacities": {"Zaloha": {"minimum": 0}}},
            {"name": "R2", "capacities": {"Zaloha": {"minimum": 0}}},
        ],
        "capacities": {},
    },
    {"name": "C", "rooms": [{"name": "R3", "capacities": {"Zaloha": {"minimum": 0}}}], "capacities": {}},
]


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ROSTERING_BUILDINGS_CONFIG_PATH", str(tmp_path / "buildings-config.yaml"))
    from rostering.persistence import config_store as config_store_module

    importlib.reload(config_store_module)
    return Workspace(root=tmp_path / "seasons")


def _helper(helper_id: int, name: str, **extra) -> dict:
    return {
        "id": helper_id,
        "name": name,
        "role_preferences": {},
        "building_preferences": [],
        "friends": [],
        "can_bring_notebook": False,
        "can_bring_camera": False,
        "unresolved_friend_names": [],
        **extra,
    }


def _seed(workspace: Workspace) -> None:
    state = workspace.load()
    state["helpers"] = [_helper(1, "Anna"), _helper(2, "Petr"), _helper(3, "Jana")]
    workspace.save(state)
    mutations.put_config(workspace, CONFIG)


def _create(workspace, name, email=None) -> int:
    """Create an Organizer by hand; returns their id."""
    return mutations.add_organizer(workspace, name, email)["organizers"][-1]["id"]


def _org(state, organizer_id) -> dict:
    return next(o for o in state["organizers"] if o["id"] == organizer_id)


def _slots(state, organizer_id=None) -> list[tuple]:
    """(role, building, room, organizer id) of the slot entries holding an Organizer."""
    return [
        (s["role"], s["building"], s["room"], s["organizer_id"])
        for s in state["manual_roles"]["structural"]
        if s.get("organizer_id") is not None and organizer_id in (None, s["organizer_id"])
    ]


def _placement(state, organizer_id):
    org = _org(state, organizer_id)
    return org.get("building"), org.get("room")


# -- created by hand ---------------------------------------------------------------


def test_an_organizer_is_created_by_hand_with_a_name_only(workspace):
    _season(workspace, "2026-jaro", ("Anna", "a@example.test"), ("Petr", "p@example.test"), ("Jana", "j@example.test"))
    state = mutations.add_organizer(workspace, "  Marie Nová ")

    [org] = state["organizers"]
    assert org["name"] == "Marie Nová"
    assert isinstance(org["id"], int)
    assert org["person_id"]
    assert org["email"] is None
    assert _placement(state, org["id"]) == (None, None)  # no placement without a slot
    assert len(state["helpers"]) == 3  # not a Helper
    assert Workspace(root=workspace.root).load()["organizers"] == state["organizers"]  # saved per Season


def test_an_organizer_needs_a_name(workspace):
    with pytest.raises(mutations.RosteringError):
        mutations.add_organizer(workspace, "   ")


def test_an_optional_email_is_stored_normalized(workspace):
    state = mutations.add_organizer(workspace, "Marie", "  Marie@Example.TEST ")
    assert state["organizers"][0]["email"] == "marie@example.test"


def test_organizer_ids_are_never_reused(workspace):
    first = _create(workspace, "A")
    second = _create(workspace, "B")
    mutations.delete_organizer(workspace, second)
    third = _create(workspace, "C")

    assert len({first, second, third}) == 3
    assert third > second


def test_organizers_round_trip_through_serialization():
    org = Organizer(id=4, name="Marie", person_id="p1", email="m@example.test", building="B", room="R1")
    assert organizer_from_dict(organizer_to_dict(org)) == org
    assert organizer_from_dict({"id": 1, "name": "X"}) == Organizer(id=1, name="X")


def test_an_organizer_can_be_renamed_and_given_an_email(workspace):
    org_id = _create(workspace, "Marie")
    state = mutations.update_organizer(workspace, org_id, name="Marie Nová", email="M@Example.test")
    assert (_org(state, org_id)["name"], _org(state, org_id)["email"]) == ("Marie Nová", "m@example.test")
    state = mutations.update_organizer(workspace, org_id, email="")
    assert _org(state, org_id)["email"] is None


# -- slots set the placement --------------------------------------------------------


def test_a_room_scoped_slot_sets_building_and_room(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")

    state = mutations.assign_organizer(workspace, org_id, "VedouciMistnosti", "B", "R1")

    assert _placement(state, org_id) == ("B", "R1")
    assert _slots(state) == [("VedouciMistnosti", "B", "R1", org_id)]


def test_a_building_scoped_slot_sets_the_building(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")

    for role in ("VedouciBudovy", "TechnickaPodpora"):
        state = mutations.assign_organizer(workspace, org_id, role, "C")
        assert _placement(state, org_id) == ("C", None)


def test_pravá_ruka_takes_a_building_or_a_room(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")

    assert _placement(mutations.assign_organizer(workspace, org_id, "PravaRuka", "B"), org_id) == ("B", None)
    assert _placement(mutations.assign_organizer(workspace, org_id, "PravaRuka", "B", "R2"), org_id) == ("B", "R2")


def test_a_scope_mismatch_is_a_structural_error_and_changes_nothing(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    before = mutations.get_state(workspace)

    for role, building, room in [
        ("VedouciBudovy", "B", "R1"),  # a building-scoped slot takes no Room
        ("TechnickaPodpora", "B", "R1"),
        ("VedouciMistnosti", "B", None),  # a room-scoped slot needs one
        ("VedouciMistnosti", "B", "R3"),  # R3 is not in B
        ("VedouciBudovy", "Nowhere", None),
        ("Registrace", "B", None),  # an Additional role is Helper-only
        ("Nonsense", "B", None),
    ]:
        with pytest.raises(mutations.RosteringError):
            mutations.assign_organizer(workspace, org_id, role, building, room)
    with pytest.raises(mutations.RosteringError):
        mutations.assign_organizer(workspace, 999, "VedouciBudovy", "B")

    assert mutations.get_state(workspace) == before


def test_assigning_to_a_new_slot_moves_the_placement_and_removes_previous_entries(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, org_id, "VedouciMistnosti", "B", "R1")

    state = mutations.assign_organizer(workspace, org_id, "TechnickaPodpora", "C")

    assert _placement(state, org_id) == ("C", None)
    assert _slots(state) == [("TechnickaPodpora", "C", None, org_id)]


def test_slots_at_the_same_placement_coexist(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, org_id, "VedouciBudovy", "B")

    state = mutations.assign_organizer(workspace, org_id, "TechnickaPodpora", "B")

    assert _placement(state, org_id) == ("B", None)
    assert {s[0] for s in _slots(state)} == {"VedouciBudovy", "TechnickaPodpora"}


def test_removing_them_from_every_slot_clears_the_placement(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, org_id, "VedouciBudovy", "B")
    mutations.assign_organizer(workspace, org_id, "TechnickaPodpora", "B")

    state = mutations.unassign_organizer(workspace, org_id, "VedouciBudovy", "B")
    assert _placement(state, org_id) == ("B", None)  # still holds one slot
    state = mutations.unassign_organizer(workspace, org_id, "TechnickaPodpora", "B")

    assert _placement(state, org_id) == (None, None)
    assert _slots(state) == []
    assert state["organizers"]  # the Organizer stays tracked


@pytest.mark.parametrize(
    "role, room",
    [("VedouciBudovy", None), ("PravaRuka", None), ("VedouciMistnosti", "R1"), ("TechnickaPodpora", None)],
)
def test_every_slot_can_have_several_holders(workspace, role, room):
    _seed(workspace)
    first, second = _create(workspace, "Marie"), _create(workspace, "Karel")
    mutations.assign_organizer(workspace, first, role, "B", room)

    state = mutations.assign_organizer(workspace, second, role, "B", room)

    assert {s[3] for s in _slots(state)} == {first, second}
    assert _placement(state, first) == _placement(state, second) == ("B", room)


def test_deleting_an_organizer_asks_when_they_hold_a_slot_and_then_frees_it(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, org_id, "VedouciBudovy", "B")

    with pytest.raises(mutations.ConfirmationRequired):
        mutations.delete_organizer(workspace, org_id)
    state = mutations.delete_organizer(workspace, org_id, confirmed=True)

    assert state["organizers"] == []
    assert state["manual_roles"]["structural"] == []


def test_slot_holders_are_picked_or_created_by_name_for_a_grid_cell(workspace):
    _seed(workspace)
    marie = _create(workspace, "Marie Nová")

    # Picking an existing Organizer by name (case and diacritics ignored) ...
    state = mutations.set_slot_holders(workspace, "VedouciBudovy", "B", None, ["marie nova"])
    assert _slots(state) == [("VedouciBudovy", "B", None, marie)]
    # ... and creating one on the spot for a name nobody tracks.
    state = mutations.set_slot_holders(workspace, "TechnickaPodpora", "C", None, ["Karel Nový"])
    [karel] = [o for o in state["organizers"] if o["name"] == "Karel Nový"]
    assert _slots(state, karel["id"]) == [("TechnickaPodpora", "C", None, karel["id"])]
    assert (karel["building"], karel["room"]) == ("C", None)
    # Emptying the cell removes the holder and clears their placement.
    state = mutations.set_slot_holders(workspace, "VedouciBudovy", "B", None, [])
    assert _placement(state, marie) == (None, None)


def test_typing_a_helpers_name_into_a_slot_creates_an_organizer_not_a_helper_reference(workspace):
    _seed(workspace)

    state = mutations.set_slot_holders(workspace, "VedouciBudovy", "B", None, ["Anna"])

    [entry] = state["manual_roles"]["structural"]
    assert entry["organizer_id"] is not None
    assert entry.get("helper_id") is None and entry.get("helper_name") is None
    assert [h["name"] for h in state["helpers"]] == ["Anna", "Petr", "Jana"]  # the Helper is untouched


def test_the_raw_manual_role_writer_rejects_an_unknown_organizer_and_a_scope_mismatch(workspace):
    _seed(workspace)
    org_id = _create(workspace, "Marie")
    bad = [
        {"role": "VedouciBudovy", "building": "B", "room": None, "organizer_id": 999},
        {"role": "VedouciMistnosti", "building": "B", "room": None, "organizer_id": org_id},
    ]
    for entry in bad:
        with pytest.raises(mutations.RosteringError):
            mutations.put_manual_roles(workspace, {"structural": [entry], "overlay": []})


# -- no solver Role, no capacity ------------------------------------------------------


def test_an_organizer_never_takes_a_solved_role_or_a_role_capacity(workspace):
    _seed(workspace)
    marie = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, marie, "VedouciMistnosti", "B", "R1")

    state = mutations.solve(workspace)

    assert {a["helper_id"] for a in state["assignments"]} == {1, 2, 3}
    assert {a["helper_name"] for a in state["assignments"]} == {"Anna", "Petr", "Jana"}


def test_the_solver_ignores_organizers_on_the_competition():
    room = Room(name="R1", capacities={Role.Opravovatel: RoleCapacity(1), Role.Zaloha: RoleCapacity(0)})
    buildings = {"B": Building(name="B", rooms=[room])}
    helpers = [Helper(id=1, name="Anna"), Helper(id=2, name="Petr")]
    organizers = [Organizer(id=1, name="Marie", building="B", room="R1")]

    plain = solve_competition(Competition(buildings=buildings, helpers=helpers), SolverConfig())
    with_organizers = solve_competition(
        Competition(buildings=buildings, helpers=helpers, organizers=organizers), SolverConfig()
    )

    def key(result):
        return sorted((a.helper_id, a.building, a.room, a.role.name) for a in result.assignments)

    assert key(with_organizers) == key(plain)
    assert len(with_organizers.assignments) == 2


def test_hand_placing_an_organizer_is_never_blocked_and_breaks_no_rule_by_itself(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    broken_before = mutations.broken_rules(mutations.get_state(workspace))
    marie = _create(workspace, "Marie")

    # Anywhere valid, however many others are already there.
    for room in ("R1", "R2", "R1"):
        state = mutations.assign_organizer(workspace, marie, "VedouciMistnosti", "B", room)

    assert _placement(state, marie) == ("B", "R1")
    assert mutations.broken_rules(state) == broken_before


# -- saved per Season and in Versions ------------------------------------------------


def test_versions_snapshot_and_restore_organizers_and_their_placement(workspace):
    _season(workspace, "2026-jaro", ("Anna", "anna@example.test"))
    mutations.put_config(workspace, CONFIG)
    marie = _create(workspace, "Marie")
    mutations.assign_organizer(workspace, marie, "VedouciMistnosti", "B", "R1")
    saved = mutations.save_version(workspace, "with Marie")

    mutations.assign_organizer(workspace, marie, "TechnickaPodpora", "C")
    mutations.add_organizer(workspace, "Karel")
    restored = mutations.restore_version(workspace, saved["slug"])

    assert [o["name"] for o in restored["organizers"]] == ["Marie"]
    assert _placement(restored, marie) == ("B", "R1")
    assert _slots(restored) == [("VedouciMistnosti", "B", "R1", marie)]


# -- legacy entries -----------------------------------------------------------------


def _legacy(workspace, structural):
    return mutations.put_manual_roles(workspace, {"structural": structural, "overlay": []})


def test_legacy_helper_id_and_typed_entries_stay_readable_and_are_marked_untracked(workspace):
    _seed(workspace)
    state = _legacy(
        workspace,
        [
            {"role": "VedouciBudovy", "building": "B", "room": None, "helper_id": 1, "helper_name": None},
            {"role": "TechnickaPodpora", "building": "C", "room": None, "helper_id": None, "helper_name": "Pan Domovník"},
        ],
    )

    assert [s["helper_id"] for s in state["manual_roles"]["structural"]] == [1, None]
    assert mutations.legacy_slot_entries(state) == [
        {"role": "VedouciBudovy", "building": "B", "room": None, "name": "Anna"},
        {"role": "TechnickaPodpora", "building": "C", "room": None, "name": "Pan Domovník"},
    ]
    assert state["organizers"] == []  # nothing was silently promoted


def test_legacy_entries_still_export_until_replaced(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    _legacy(
        workspace,
        [
            {"role": "VedouciBudovy", "building": "B", "room": None, "helper_id": None, "helper_name": "Pan Vedoucí"},
            {"role": "TechnickaPodpora", "building": "C", "room": None, "helper_id": 1, "helper_name": None},
        ],
    )

    values = _sheet_values(mutations.export_xlsx_bytes(workspace))

    assert "Pan Vedoucí" in values
    assert any(v.startswith("Anna") for v in values)


def test_a_legacy_entry_is_replaced_once_it_is_no_longer_named(workspace):
    _seed(workspace)
    _legacy(
        workspace,
        [{"role": "VedouciBudovy", "building": "B", "room": None, "helper_id": None, "helper_name": "Pan Vedoucí"}],
    )

    state = mutations.set_slot_holders(workspace, "VedouciBudovy", "B", None, ["Marie"])

    [entry] = state["manual_roles"]["structural"]
    assert _org(state, entry["organizer_id"])["name"] == "Marie"
    assert mutations.legacy_slot_entries(state) == []


def test_a_legacy_entry_can_be_removed_from_its_cell(workspace):
    _seed(workspace)
    _legacy(
        workspace,
        [{"role": "TechnickaPodpora", "building": "B", "room": None, "helper_id": None, "helper_name": "Pan Domovník"}],
    )

    state = mutations.set_slot_holders(workspace, "TechnickaPodpora", "B", None, [])

    assert state["manual_roles"]["structural"] == []


def test_a_legacy_entry_is_kept_when_another_organizer_is_added_to_a_multi_holder_cell(workspace):
    _seed(workspace)
    _legacy(
        workspace,
        [{"role": "TechnickaPodpora", "building": "B", "room": None, "helper_id": None, "helper_name": "Pan Domovník"}],
    )

    state = mutations.set_slot_holders(workspace, "TechnickaPodpora", "B", None, ["Pan Domovník", "Marie"])

    assert [s.get("helper_name") for s in state["manual_roles"]["structural"]] == ["Pan Domovník", None]
    assert len(state["organizers"]) == 1


# -- export -------------------------------------------------------------------------


def _sheet_values(xlsx: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    return [str(c.value) for row in wb["Pomocníci v místnostech"].iter_rows() for c in row if c.value]


def _row_cells(xlsx: bytes, label: str) -> list[str]:
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))["Pomocníci v místnostech"]
    for row in ws.iter_rows():
        if row[0].value == label:
            return [str(c.value) for c in row[1:] if c.value]
    raise AssertionError(f"no row {label!r}")


def test_the_exported_slot_cells_show_organizer_names(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    marie = _create(workspace, "Marie Nová")
    karel = _create(workspace, "Karel Nový")
    jitka = _create(workspace, "Jitka Nová")
    mutations.assign_organizer(workspace, marie, "VedouciBudovy", "B")
    mutations.assign_organizer(workspace, karel, "VedouciMistnosti", "B", "R2")
    mutations.assign_organizer(workspace, jitka, "TechnickaPodpora", "C")

    xlsx = mutations.export_xlsx_bytes(workspace)

    assert _row_cells(xlsx, "Vedoucí budovy") == ["Marie Nová"]
    assert _row_cells(xlsx, "Vedoucí místností") == ["Karel Nový"]
    assert _row_cells(xlsx, "Technická podpora") == ["Jitka Nová"]


def test_an_organizer_is_counted_once_in_the_building_they_are_placed_in(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    marie = _create(workspace, "Marie Nová")
    mutations.assign_organizer(workspace, marie, "VedouciBudovy", "C")

    wb = openpyxl.load_workbook(io.BytesIO(mutations.export_xlsx_bytes(workspace)))

    names_in_c = [str(c.value) for row in wb["C"].iter_rows() for c in row if c.value]
    assert any("Marie Nová" in v for v in names_in_c)


# -- recognition across Seasons ---------------------------------------------------------


def _survey(rows) -> bytes:
    buffer = io.BytesIO()
    pd.DataFrame(
        {
            "Časová značka": [datetime(2026, 2, 1)] * len(rows),
            "Tvoje jméno a příjmení": [name for name, _ in rows],
            "E-mailová adresa": [email for _, email in rows],
        }
    ).to_excel(buffer, index=False)
    return buffer.getvalue()


def _season(workspace, label, *rows):
    mutations.new_season(workspace)
    return mutations.upload_responses(workspace, _survey(rows), "s.xlsx", label=label)


def _open(workspace, label):
    ids = {s["label"]: s["id"] for s in mutations.list_seasons(workspace)}
    return mutations.open_season(workspace, ids[label])


def test_an_email_match_links_a_new_organizer_to_the_person_an_earlier_season_recorded(workspace):
    earlier = _season(workspace, "2025-podzim", ("Anna Nováková", "anna@example.test"))
    _season(workspace, "2026-jaro", ("Petr", "petr@example.test"))

    state = mutations.add_organizer(workspace, "Aňa", "ANNA@example.test")

    assert state["organizers"][0]["person_id"] == earlier["helpers"][0]["person_id"]


def test_an_organizer_without_an_email_is_only_an_uncertain_name_match(workspace):
    earlier = _season(workspace, "2025-podzim", ("Anna Nováková", "anna@example.test"))
    _season(workspace, "2026-jaro", ("Petr", "petr@example.test"))

    state = mutations.add_organizer(workspace, "anna novakova")
    org_id = state["organizers"][0]["id"]

    assert state["organizers"][0]["person_id"] != earlier["helpers"][0]["person_id"]  # nothing auto-linked
    [entry] = mutations.get_uncertain_organizer_matches(workspace)
    assert entry["organizer_id"] == org_id and entry["organizer_name"] == "anna novakova"
    [candidate] = entry["candidates"]
    assert candidate["person_id"] == earlier["helpers"][0]["person_id"]
    assert candidate["season"] == "2025-podzim"

    linked = mutations.link_organizer(workspace, org_id, candidate["person_id"])
    assert _org(linked, org_id)["person_id"] == candidate["person_id"]
    assert mutations.get_uncertain_organizer_matches(workspace) == []


def test_entering_an_email_later_enables_the_confident_match(workspace):
    earlier = _season(workspace, "2025-podzim", ("Anna Nováková", "anna@example.test"))
    _season(workspace, "2026-jaro", ("Petr", "petr@example.test"))
    org_id = mutations.add_organizer(workspace, "Anna")["organizers"][0]["id"]

    state = mutations.update_organizer(workspace, org_id, email="anna@example.test")

    assert _org(state, org_id)["person_id"] == earlier["helpers"][0]["person_id"]


def test_a_rejected_organizer_pairing_is_not_proposed_again(workspace):
    earlier = _season(workspace, "2025-podzim", ("Anna Nováková", "anna@example.test"))
    _season(workspace, "2026-jaro", ("Petr", "petr@example.test"))
    org_id = mutations.add_organizer(workspace, "Anna Nováková")["organizers"][0]["id"]

    mutations.reject_organizer_match(workspace, org_id, earlier["helpers"][0]["person_id"])

    assert mutations.get_uncertain_organizer_matches(workspace) == []


def test_someone_known_as_an_organizer_who_registers_later_is_offered_a_link_never_promoted(workspace):
    _season(workspace, "2025-podzim", ("Petr", "petr@example.test"))
    mutations.put_config(workspace, CONFIG)
    marie = mutations.add_organizer(workspace, "Marie Nová")["organizers"][0]
    mutations.assign_organizer(workspace, marie["id"], "VedouciBudovy", "B")

    state = _season(workspace, "2026-jaro", ("Marie Nová", "marie@example.test"))

    [entry] = mutations.get_uncertain_matches(workspace)
    assert entry["helper_name"] == "Marie Nová"
    assert [c["person_id"] for c in entry["candidates"]] == [marie["person_id"]]
    assert state["organizers"] == []  # this Season has no Organizers: nobody was promoted
    assert [h["name"] for h in state["helpers"]] == ["Marie Nová"]

    # Confirming the link points the Helper at the Organizer's Person; still a Helper.
    linked = mutations.link_helper(workspace, state["helpers"][0]["id"], marie["person_id"])
    assert linked["helpers"][0]["person_id"] == marie["person_id"]
    assert _open(workspace, "2025-podzim")["organizers"][0]["id"] == marie["id"]


def test_a_confident_email_match_links_a_helper_to_a_known_organizers_person(workspace):
    _season(workspace, "2025-podzim", ("Petr", "petr@example.test"))
    marie = mutations.add_organizer(workspace, "Marie Nová", "marie@example.test")["organizers"][0]

    state = _season(workspace, "2026-jaro", ("M. Nová", "MARIE@example.test"))

    assert state["helpers"][0]["person_id"] == marie["person_id"]
    assert state["organizers"] == []  # linked to the Person, not promoted
    assert [p["person_id"] for p in mutations.list_persons(workspace) if "marie@example.test" in p["emails"]] == [
        marie["person_id"]
    ]


def test_deleting_a_season_forgets_the_organizers_only_it_recorded(workspace):
    _season(workspace, "2025-podzim", ("Petr", "petr@example.test"))
    marie = mutations.add_organizer(workspace, "Marie Nová", "marie@example.test")["organizers"][0]
    _season(workspace, "2026-jaro", ("Anna", "anna@example.test"))
    assert marie["person_id"] in {p["person_id"] for p in mutations.list_persons(workspace)}

    ids = {s["label"]: s["id"] for s in mutations.list_seasons(workspace)}
    mutations.delete_season(workspace, ids["2025-podzim"])

    assert marie["person_id"] not in {p["person_id"] for p in mutations.list_persons(workspace)}


def test_organizer_records_do_not_clash_with_helper_ids_in_person_links(workspace):
    state = _season(workspace, "2025-podzim", ("Petr", "petr@example.test"), ("Anna", "anna@example.test"))
    helper_ids = {h["id"] for h in state["helpers"]}
    org_id = mutations.add_organizer(workspace, "Karel")["organizers"][0]["id"]
    assert org_id in helper_ids  # the id spaces are separate and may overlap

    assert mutations.get_person_links(workspace) == {}
    assert mutations.get_uncertain_matches(workspace) == []
    assert mutations.get_uncertain_organizer_matches(workspace) == []
