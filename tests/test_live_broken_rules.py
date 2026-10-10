"""The live Broken-rule check: current Assignments judged against the current
rules on every render, the before/after diff behind the drop toast, and the
agreement between the checker and the solver's own bent rules.

Mutation-layer tests run against a temp-dir workspace seeded with synthetic
data (never anything from data/); the agreement test drives the solver on
domain objects.
"""
import io
import random

import openpyxl
import pytest

from rostering.domain import (
    Assignment,
    Building,
    Competition,
    Helper,
    Role,
    RoleCapacity,
    Room,
    RuleInstance,
)
from rostering.persistence.workspace import Workspace
from rostering.solver import rules
from rostering.solver.checker import check_roster, newly_broken
from rostering.solver.model import SolverConfig, solve_competition
from rostering.solver.rules import Relaxation, RuleFamily, Tier
from rostering.webapp import mutations


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    return Workspace(root=tmp_path / "workspace")


def _helper(hid, name, camera=False):
    return {
        "id": hid,
        "name": name,
        "role_preferences": {},
        "building_preferences": [],
        "friends": [],
        "can_bring_notebook": False,
        "can_bring_camera": camera,
        "unresolved_friend_names": [],
    }


def _assignment(hid, name, building, room, role):
    return {"helper_id": hid, "helper_name": name, "building": building, "room": room, "role": role}


def _seed(workspace, config, helpers, assignments):
    mutations.put_config(workspace, config)
    state = workspace.load()
    state["helpers"] = helpers
    state["assignments"] = assignments
    workspace.save(state)
    return workspace.load()


def _room(name, **minimums):
    return {"name": name, "capacities": {role: {"minimum": n} for role, n in minimums.items()}}


TWO_ROOMS = [
    {
        "name": "B",
        "rooms": [_room("R1", Fotograf=1), _room("R2")],
        "capacities": {"Skenovac": {"minimum": 2}},
    }
]


def _by_kind(broken):
    return {b.instance.kind: b for b in broken}


def _seed_roster(workspace):
    """Anna (camera) is Fotograf in R1, Petr and Eva scan in R2 — every rule holds."""
    helpers = [_helper(1, "Anna", camera=True), _helper(2, "Petr"), _helper(3, "Eva")]
    assignments = [
        _assignment(1, "Anna", "B", "R1", "Fotograf"),
        _assignment(2, "Petr", "B", "R2", "Skenovac"),
        _assignment(3, "Eva", "B", "R2", "Skenovac"),
    ]
    return _seed(workspace, TWO_ROOMS, helpers, assignments)


# -- the live check ----------------------------------------------------------


def test_a_roster_that_keeps_every_rule_has_no_broken_rules(workspace):
    state = _seed_roster(workspace)

    assert mutations.broken_rules(state) == []


def test_no_roster_yet_is_not_reported_as_broken(workspace):
    state = _seed(workspace, TWO_ROOMS, [_helper(1, "Anna", camera=True)], [])

    assert mutations.broken_rules(state) == []


def test_a_hand_move_that_breaks_a_minimum_shows_it_at_once_and_moving_back_clears_it(workspace):
    _seed_roster(workspace)

    broken = mutations.broken_rules(mutations.move_helper(workspace, 1, "B", "R1", "Zaloha"))

    instance = RuleInstance("room_exact", ("B", "R1", "Fotograf"))
    assert [b.instance for b in broken] == [instance]
    assert broken[0].family == "minimums"
    assert broken[0].amount == 1
    assert broken[0].line == "Místnost R1 · Fotograf: 0 z 1 požadovaných (chybí 1)"

    fixed = mutations.move_helper(workspace, 1, "B", "R1", "Fotograf")
    assert mutations.broken_rules(fixed) == []


def test_a_building_limit_shortfall_is_a_count(workspace):
    _seed_roster(workspace)

    broken = mutations.broken_rules(mutations.move_helper(workspace, 3, "B", "R2", "Zaloha"))

    assert [b.instance for b in broken] == [RuleInstance("building_exact", ("B", "Skenovac"))]
    assert broken[0].line == "Budova B · Skenovač: 1 z 2 požadovaných (chybí 1)"


def test_an_equipment_violation_names_the_helper(workspace):
    _seed_roster(workspace)

    broken = mutations.broken_rules(mutations.move_helper(workspace, 2, "B", "R1", "Fotograf"))

    equipment = _by_kind(broken)["equipment"]
    assert equipment.instance == RuleInstance("equipment", (2,))
    assert equipment.family == "equipment"
    assert equipment.line == "Pomocník Petr nemá fotoaparát, ale je Fotograf"


def test_broken_rules_come_in_tier_order_minimums_first(workspace):
    _seed_roster(workspace)

    # Petr (no camera) becomes the second Fotograf (R1 holds exactly one) while
    # R2 loses a scanner (the Building holds exactly two).
    mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")
    broken = mutations.broken_rules(mutations.get_state(workspace))

    assert [b.family for b in broken] == ["minimums", "minimums", "equipment"]


def test_a_building_limit_overshoot_is_a_count_too(workspace):
    _seed_roster(workspace)

    broken = mutations.broken_rules(mutations.move_helper(workspace, 1, "B", "R1", "Skenovac"))

    # Three scan where the Building's limit is exactly two; Anna is also
    # missing from R1's Fotograf minimum, which is a separate rule.
    building = next(b for b in broken if b.instance.kind == "building_exact")
    assert building.instance == RuleInstance("building_exact", ("B", "Skenovac"))
    assert building.amount == 1
    assert building.line == "Budova B · Skenovač: 3 z 2 požadovaných (přebývá 1)"


def test_stale_assignments_to_removed_rooms_do_not_crash_the_check(workspace):
    state = _seed_roster(workspace)
    state["assignments"][2] = _assignment(3, "Eva", "Gone", "Nowhere", "Skenovac")

    # Eva no longer counts anywhere: the shortfall, not an exception.
    assert [b.instance.kind for b in mutations.broken_rules(state)] == ["building_exact"]


def test_each_instance_carries_the_cells_and_chips_it_affects(workspace):
    _seed_roster(workspace)
    state = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")
    state = mutations.move_helper(workspace, 3, "B", "R2", "Zaloha")

    broken = _by_kind(mutations.broken_rules(state))

    assert broken["building_exact"].cells == (("B", "R1", "Skenovac"), ("B", "R2", "Skenovac"))
    assert broken["equipment"].helper_ids == (2,)
    marks = mutations.broken_rule_marks(mutations.broken_rules(state))
    assert marks["helpers"] == [{"helper_id": 2, "line": "Pomocník Petr nemá fotoaparát, ale je Fotograf"}]
    # The equipment rule marks Petr's whole Room (R1); the counts their cells.
    assert {(c["building"], c["room"], c["role"]) for c in marks["cells"]} == {
        ("B", "R1", "Fotograf"),
        ("B", "R1", "Skenovac"),
        ("B", "R2", "Skenovac"),
        ("B", "R1", None),
    }


def test_each_instance_has_a_go_fix_target(workspace):
    _seed_roster(workspace)
    mutations.move_helper(workspace, 1, "B", "R1", "Zaloha")
    state = mutations.move_helper(workspace, 2, "B", "R2", "Fotograf")

    broken = _by_kind(mutations.broken_rules(state))

    assert broken["room_exact"].fix.tab == "buildings"
    assert (broken["room_exact"].fix.building, broken["room_exact"].fix.room) == ("B", "R1")
    assert broken["building_exact"].fix.tab == "buildings"
    assert broken["building_exact"].fix.building == "B"
    assert broken["equipment"].fix.tab == "helpers"
    assert broken["equipment"].fix.helper_id == 2


def test_nothing_about_broken_rules_is_persisted_by_a_hand_move(workspace):
    before = _seed_roster(workspace)

    mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")

    after = workspace.load()
    assert set(after) == set(before)
    assert set(after["diagnostics"]) == set(before["diagnostics"])


# -- the before/after diff and the drop toast ---------------------------------


def test_the_diff_reports_only_instances_new_in_the_after_state(workspace):
    before = _seed_roster(workspace)
    mid = mutations.move_helper(workspace, 1, "B", "R1", "Zaloha")  # breaks the R1 minimum
    after = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")  # fixes it, breaks camera

    assert [b.instance for b in mutations.newly_broken_rules(before, mid)] == [
        RuleInstance("room_exact", ("B", "R1", "Fotograf"))
    ]
    diff = mutations.newly_broken_rules(mid, after)
    assert RuleInstance("equipment", (2,)) in [b.instance for b in diff]
    assert RuleInstance("room_exact", ("B", "R1", "Fotograf")) not in [b.instance for b in diff]


def test_a_move_that_breaks_equipment_toasts_a_line_worded_like_the_banner(workspace):
    before = _seed_roster(workspace)
    after = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")

    lines = mutations.move_toast_lines(before, after)

    banner = [b.line for b in mutations.broken_rules(after) if b.family == "equipment"]
    assert lines == banner == ["Pomocník Petr nemá fotoaparát, ale je Fotograf"]


def test_minimums_never_toast(workspace):
    before = _seed_roster(workspace)
    after = mutations.move_helper(workspace, 1, "B", "R1", "Zaloha")

    assert mutations.broken_rules(after)  # the banner shows it ...
    assert mutations.move_toast_lines(before, after) == []  # ... the toast does not


def test_a_rule_already_broken_before_the_drop_does_not_toast_again(workspace):
    _seed_roster(workspace)
    before = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")  # Petr breaks camera
    after = mutations.move_helper(workspace, 2, "B", "R2", "Fotograf")  # still Fotograf, elsewhere

    assert RuleInstance("equipment", (2,)) in [b.instance for b in mutations.broken_rules(after)]
    assert mutations.move_toast_lines(before, after) == []


def test_a_move_that_fixes_or_changes_nothing_toasts_nothing(workspace):
    _seed_roster(workspace)
    broken_state = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")
    fixed = mutations.move_helper(workspace, 2, "B", "R2", "Skenovac")
    same = mutations.move_helper(workspace, 2, "B", "R2", "Skenovac")

    assert mutations.move_toast_lines(broken_state, fixed) == []
    assert mutations.move_toast_lines(fixed, same) == []


def test_a_hand_move_is_never_refused_whatever_it_breaks(workspace):
    _seed_roster(workspace)

    state = mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")

    moved = next(a for a in state["assignments"] if a["helper_id"] == 2)
    assert (moved["room"], moved["role"]) == ("R1", "Fotograf")


def test_export_is_not_blocked_by_broken_rules_and_shows_nothing_about_them(workspace):
    _seed_roster(workspace)
    clean = mutations.export_xlsx_bytes(workspace)
    mutations.move_helper(workspace, 2, "B", "R1", "Fotograf")
    mutations.move_helper(workspace, 3, "B", "R2", "Zaloha")
    assert mutations.broken_rules(mutations.get_state(workspace))

    exported = mutations.export_xlsx_bytes(workspace)

    assert exported[:2] == b"PK"
    workbook = openpyxl.load_workbook(io.BytesIO(exported))
    text = " ".join(
        str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value is not None
    )
    assert "Petr" in text
    assert "required" not in text and "no camera" not in text
    # Same sheets as the clean roster's export, broken rules or not.
    assert workbook.sheetnames == openpyxl.load_workbook(io.BytesIO(clean)).sheetnames


# -- the checker on domain objects: the extension point -----------------------


def _config():
    return SolverConfig(time_limit_seconds=5)


def _domain_room(name, **minimums):
    caps = {Role.Zaloha: RoleCapacity(minimum=0)}
    for role_name, minimum in minimums.items():
        caps[Role[role_name]] = RoleCapacity(minimum=minimum)
    return Room(name=name, capacities=caps)


def _projection(broken_rules):
    return {b.instance: (b.family, b.amount, b.line) for b in broken_rules}


def _random_competition(rng):
    buildings = {}
    for b in range(rng.randint(1, 2)):
        name = f"B{b}"
        rooms = []
        for r in range(rng.randint(1, 3)):
            minimums = {role.name: rng.randint(0, 3) for role in (Role.Fotograf, Role.Skenovac, Role.Menic)}
            rooms.append(_domain_room(f"B{b}R{r}", **{k: v for k, v in minimums.items() if v}))
        capacities = {}
        if rng.random() < 0.5:
            capacities[Role.Opravovatel] = RoleCapacity(minimum=rng.randint(1, 4))
        buildings[name] = Building(name=name, rooms=rooms, capacities=capacities)
    helpers = [
        Helper(id=i, name=f"H{i}", can_bring_camera=rng.random() < 0.3) for i in range(1, rng.randint(2, 9))
    ]
    return Competition(buildings=buildings, helpers=helpers)


@pytest.mark.parametrize("seed", range(12))
def test_the_checker_agrees_with_the_solvers_own_bent_rules(seed):
    comp = _random_competition(random.Random(seed))

    result = solve_competition(comp, _config())

    assert _projection(check_roster(comp, result.assignments)) == _projection(result.broken_rules)


def test_the_checker_agrees_with_the_solver_when_equipment_is_forced_to_bend():
    building = Building(name="B", rooms=[_domain_room("R1")])
    comp = Competition(buildings={"B": building}, helpers=[Helper(id=1, name="Nocam", can_bring_camera=False)])
    fixed = Assignment(helper_id=1, helper_name="Nocam", building="B", room="R1", role=Role.Fotograf)

    result = solve_competition(comp, _config(), fixed_assignments=[fixed])

    assert result.broken_rules
    assert _projection(check_roster(comp, result.assignments)) == _projection(result.broken_rules)


def _must_be_fotograf(helper_id):
    """A stand-in for a later rule family, stated once for the solver and once
    for the live checker."""

    def relax(ctx):
        return [
            Relaxation(
                instance=RuleInstance("must_be_fotograf", (helper_id,)),
                slack=1 - ctx.assign_role[helper_id, Role.Fotograf],
                max_units=1,
                describe=lambda _n: f"Helper {helper_id} must be Fotograf",
            )
        ]

    def check(ctx):
        from rostering.domain import BrokenRule

        role = next(a.role for a in ctx.assignments if a.helper_id == helper_id)
        if role == Role.Fotograf:
            return []
        return [
            BrokenRule(
                instance=RuleInstance("must_be_fotograf", (helper_id,)),
                family="must_be_fotograf",
                amount=1,
                line=f"Helper {helper_id} must be Fotograf",
            )
        ]

    return RuleFamily("must_be_fotograf", Tier.TAG_RESTRICTIONS, relax, check=check)


def test_a_registered_rule_family_is_used_by_both_the_solver_and_the_live_checker(monkeypatch):
    monkeypatch.setattr(rules, "_registry", rules.rule_families())
    rules.register_rule_family(_must_be_fotograf(1))
    building = Building(name="B", rooms=[_domain_room("R1")])
    comp = Competition(buildings={"B": building}, helpers=[Helper(id=1, name="Nocam", can_bring_camera=False)])

    result = solve_competition(comp, _config())

    # Nobody can be Fotograf without a camera: the Tag-tier stand-in bends.
    assert [b.instance for b in result.broken_rules] == [RuleInstance("must_be_fotograf", (1,))]
    assert _projection(check_roster(comp, result.assignments)) == _projection(result.broken_rules)


def test_a_rule_family_without_a_live_check_cannot_be_registered(monkeypatch):
    monkeypatch.setattr(rules, "_registry", rules.rule_families())
    half = RuleFamily("half_way", Tier.TAG_RESTRICTIONS, lambda ctx: [])

    with pytest.raises(ValueError):
        rules.register_rule_family(half)


def test_the_diff_compares_by_rule_identity_not_by_size():
    building = Building(
        name="B",
        rooms=[_domain_room("R1", Skenovac=3)],
    )
    comp = Competition(buildings={"B": building}, helpers=[Helper(id=i, name=str(i)) for i in (1, 2)])
    scanning = lambda n: [  # noqa: E731
        Assignment(helper_id=i, helper_name=str(i), building="B", room="R1", role=Role.Skenovac) for i in range(1, n + 1)
    ]

    one_short = check_roster(comp, scanning(2))
    two_short = check_roster(comp, scanning(1))

    # Getting worse is still the same broken rule, not a newly broken one.
    assert newly_broken(one_short, two_short) == []
    assert newly_broken([], two_short) == two_short
