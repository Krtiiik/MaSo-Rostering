"""The Buildings layout read off a "Pomocníci v místnostech" sheet: Buildings and
Rooms from the merged header cells, Role counts from the coloured cells."""
import io

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from rostering.ingest.building_sheet import parse_building_sheet
from rostering.webapp import mutations

NEEDED = "FFFFE599"
GRAY = "FF999999"


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


def _put(ws, cell: str, value=None, fill=None) -> None:
    ws[cell] = value
    if fill:
        ws[cell].fill = fill if isinstance(fill, PatternFill) else _fill(fill)


def _sheet(build) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    build(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _two_buildings(ws) -> None:
    """Alfa: rooms A1, A2 and a large A3 (a merged header over D:E). Beta: B1 and
    B2, whose header is followed by a blank cell. Rows: leaders (ignored),
    Opravovatelé (3 rows), Skenovači, Fotografové, Záloha."""
    ws.merge_cells("B1:E1")
    ws.merge_cells("F1:H1")
    ws.merge_cells("D2:E2")
    _put(ws, "B1", "Alfa")
    _put(ws, "F1", "Beta")
    for cell, name in {"B2": "A1", "C2": "A2", "D2": "A3", "F2": "B1", "G2": "B2"}.items():
        _put(ws, cell, name)
    _put(ws, "A3", "Vedoucí budovy", "FFCCCCCC")
    _put(ws, "B3", "Jiří Cihelka", "FFD9D9D9")
    ws.merge_cells("A4:A6")
    _put(ws, "A4", "Opravovatelé")
    _put(ws, "A7", "Skenovači")
    _put(ws, "A8", "Fotografové")
    _put(ws, "A9", "Záloha")
    # Opravovatelé, top rows coloured and the rest gray: column B 3 of 3, C 2, D 3,
    # E 2, F 1, G 3, H 1 -- so A1 needs 3, A2 2, A3 (D and E) 5, B1 1, B2 (G, H) 4.
    for col, rows in {"B": 3, "C": 2, "D": 3, "E": 2, "F": 1, "G": 3, "H": 1}.items():
        for row in range(4, 7):
            _put(ws, f"{col}{row}", fill=NEEDED if row - 3 <= rows else GRAY)
    for col in "BCDEFGH":
        _put(ws, f"{col}7", fill=NEEDED)
    # Fotografové: one merged slot over A1 and A2 (building-wide), one over A3
    # alone (that Room's), one over all of Beta (building-wide).
    ws.merge_cells("B8:C8")
    ws.merge_cells("D8:E8")
    ws.merge_cells("F8:H8")
    for cell in ("B8", "D8", "F8"):
        _put(ws, cell, fill=NEEDED)
    for col in "BCDEFGH":
        _put(ws, f"{col}9")  # Záloha: unfilled


def _by_name(buildings: list[dict]) -> dict:
    return {b["name"]: {"caps": b["capacities"], "rooms": {r["name"]: r["capacities"] for r in b["rooms"]}} for b in buildings}


def test_buildings_and_rooms_come_from_the_header_cells():
    layout = parse_building_sheet(_sheet(_two_buildings)).buildings
    assert [b["name"] for b in layout] == ["Alfa", "Beta"]
    assert [r["name"] for r in layout[0]["rooms"]] == ["A1", "A2", "A3"]
    # B2's blank neighbour (column H) is part of that Room, not a new one.
    assert [r["name"] for r in layout[1]["rooms"]] == ["B1", "B2"]


def test_a_coloured_cell_is_one_helper_and_a_gray_or_empty_one_is_none():
    layout = _by_name(parse_building_sheet(_sheet(_two_buildings)).buildings)
    minimum = lambda caps, role: caps.get(role, {}).get("minimum", 0)  # noqa: E731
    rooms = layout["Alfa"]["rooms"]
    assert [minimum(rooms[r], "Opravovatel") for r in ("A1", "A2", "A3")] == [3, 2, 5]  # A3 is two columns wide
    assert [minimum(layout["Beta"]["rooms"][r], "Opravovatel") for r in ("B1", "B2")] == [1, 4]
    assert minimum(rooms["A1"], "Skenovac") == 1 and minimum(rooms["A3"], "Skenovac") == 2
    assert "Zaloha" not in rooms["A1"]  # an unfilled row asks for nobody


def test_a_merged_slot_counts_once_for_one_room_or_for_the_building():
    layout = _by_name(parse_building_sheet(_sheet(_two_buildings)).buildings)
    assert layout["Alfa"]["caps"] == {"Fotograf": {"minimum": 1}}  # B8:C8 over A1 and A2
    assert layout["Alfa"]["rooms"]["A3"]["Fotograf"] == {"minimum": 1}  # D8:E8 over A3 alone
    assert "Fotograf" not in layout["Alfa"]["rooms"]["A1"]
    assert layout["Beta"]["caps"] == {"Fotograf": {"minimum": 1}}  # F8:H8 over both Rooms


def test_a_theme_gray_is_not_needed_and_a_theme_accent_is():
    def build(ws) -> None:
        ws.merge_cells("B1:C1")
        _put(ws, "B1", "Alfa")
        _put(ws, "B2", "A1")
        _put(ws, "C2", "A2")
        _put(ws, "A3", "Skenovači")
        ws.merge_cells("A3:A4")
        _put(ws, "B3", fill=PatternFill("solid", fgColor=Color(theme=0, tint=-0.35)))  # darker white
        _put(ws, "B4", fill=PatternFill("solid", fgColor=Color(theme=5)))  # an accent
        _put(ws, "C3", fill=PatternFill("solid", fgColor=Color(indexed=22)))  # indexed gray
        _put(ws, "C4", fill=PatternFill("solid", fgColor=Color(indexed=13)))  # indexed yellow

    rooms = parse_building_sheet(_sheet(build)).buildings[0]["rooms"]
    assert [r["capacities"] for r in rooms] == [{"Skenovac": {"minimum": 1}}, {"Skenovac": {"minimum": 1}}]


def test_a_room_with_no_coloured_cell_is_reported():
    def build(ws) -> None:
        _put(ws, "B1", "Alfa")
        _put(ws, "B2", "A1")
        _put(ws, "C2", "A2")
        _put(ws, "A3", "Kresliči")
        _put(ws, "B3", fill=NEEDED)
        _put(ws, "C3", fill=GRAY)

    sheet = parse_building_sheet(_sheet(build))
    assert [r["capacities"] for r in sheet.buildings[0]["rooms"]] == [{"Kreslic": {"minimum": 1}}, {}]
    assert sheet.warnings == ["Místnost „A2“ (budova Alfa) nemá v tabulce žádnou barevnou buňku."]


def test_a_sheet_without_role_rows_is_refused():
    def build(ws) -> None:
        _put(ws, "B1", "Alfa")
        _put(ws, "B2", "A1")
        _put(ws, "A3", "Vedoucí budovy")

    with pytest.raises(ValueError, match="řádky rolí"):
        parse_building_sheet(_sheet(build))


def test_a_file_that_is_not_a_workbook_is_refused():
    with pytest.raises(ValueError, match="přečíst"):
        parse_building_sheet(b"not a spreadsheet")


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    from rostering.persistence.workspace import Workspace

    monkeypatch.setenv("ROSTERING_SEASONS_DIR", str(tmp_path / "seasons"))
    workspace = Workspace()
    workspace.create_season("2026-jaro")
    return workspace


def test_the_mutation_raises_a_rostering_error_instead(workspace):
    with pytest.raises(mutations.RosteringError, match="přečíst"):
        mutations.read_building_sheet(workspace, b"not a spreadsheet")
    buildings, pending, warnings = mutations.read_building_sheet(workspace, _sheet(_two_buildings))
    assert [b["name"] for b in buildings] == ["Alfa", "Beta"]
    # The fixture names Jiří Cihelka as Vedoucí budovy, and the Season has no Organizer of that name.
    assert len(warnings) == 1 and "Jiří Cihelka" in warnings[0]
    assert pending["slots"] == [] and pending["row_merges"] == []


# ---------------------------------------------------------------------- leaders and merges
def _with_leaders(ws) -> None:
    """Alfa: rooms A1, A2, A3 (columns B-D); Beta: B1, B2 (E-F). Rows: 3 Vedoucí
    budovy, 4 Pravá ruka, 5 Vedoucí místností, 6 Skenovači, 7 Fotografové, 8 Focení
    předávání cen, 9 Technická podpora."""
    ws.merge_cells("B1:D1")
    ws.merge_cells("E1:F1")
    _put(ws, "B1", "Alfa")
    _put(ws, "E1", "Beta")
    for cell, name in {"B2": "A1", "C2": "A2", "D2": "A3", "E2": "B1", "F2": "B2"}.items():
        _put(ws, cell, name)
    for row, label in enumerate(
        ["Vedoucí budovy", "Pravá ruka", "Vedoucí místností", "Skenovači", "Fotografové", "Focení předávání cen", "Technická podpora"],
        start=3,
    ):
        _put(ws, f"A{row}", label)
    ws.merge_cells("B3:D4")  # Alfa: one leader for both rows, over the whole Building
    _put(ws, "B3", "Anna Nováková")
    ws.merge_cells("E3:F3")
    _put(ws, "E3", "Bob Beran")
    ws.merge_cells("E4:F4")
    _put(ws, "E4", "Cyril Cerny, Dana Dvorakova")
    _put(ws, "B5", "Eva Ehrlichova")
    ws.merge_cells("C5:D5")
    _put(ws, "C5", "Filip Fiala")
    _put(ws, "E5", "Gita Gabrielova")
    for col in "BCDEF":
        _put(ws, f"{col}6", fill=NEEDED)
    ws.merge_cells("B7:C8")  # Fotograf over A1+A2 with Focení under it
    ws.merge_cells("E7:F7")
    for cell in ("B7", "D7", "E7"):
        _put(ws, cell, fill=NEEDED)
    _put(ws, "B9", "Hana Horka, Ivo Ivanek")
    _put(ws, "F9", "Jan Janda")


def test_leadership_cells_become_slots_at_the_building_or_the_first_room_they_cover():
    slots = parse_building_sheet(_sheet(_with_leaders)).slots
    assert [(c.role, c.building, c.room, c.names) for c in slots] == [
        ("VedouciBudovy", "Alfa", None, ["Anna Nováková"]),
        ("PravaRuka", "Alfa", None, ["Anna Nováková"]),  # filed at the Building: the cell is tall with Vedoucí budovy
        ("VedouciBudovy", "Beta", None, ["Bob Beran"]),
        ("PravaRuka", "Beta", "B1", ["Cyril Cerny", "Dana Dvorakova"]),
        ("VedouciMistnosti", "Alfa", "A1", ["Eva Ehrlichova"]),
        ("VedouciMistnosti", "Alfa", "A2", ["Filip Fiala"]),  # the cell covers A2 and A3
        ("VedouciMistnosti", "Beta", "B1", ["Gita Gabrielova"]),
        ("TechnickaPodpora", "Alfa", None, ["Hana Horka", "Ivo Ivanek"]),
        ("TechnickaPodpora", "Beta", None, ["Jan Janda"]),
    ]


def test_merged_cells_become_sideways_and_tall_merges():
    sheet = parse_building_sheet(_sheet(_with_leaders))
    assert sheet.cell_merges == {
        "PravaRuka": {"Alfa": [["A1", "A2"], ["A2", "A3"]], "Beta": [["B1", "B2"]]},
        "VedouciMistnosti": {"Alfa": [["A2", "A3"]]},
        "Fotograf": {"Alfa": [["A1", "A2"]], "Beta": [["B1", "B2"]]},
        "FoceniPredavaniCen": {"Alfa": [["A1", "A2"]]},
    }
    assert sheet.row_merges == [
        {"building": "Alfa", "room": "A1", "row": "VedouciBudovy"},
        {"building": "Alfa", "room": "A1", "row": "Fotograf"},
    ]


def test_a_tall_cell_that_cannot_exist_is_reported_and_left_out():
    def build(ws) -> None:
        _with_leaders(ws)
        ws.unmerge_cells("B3:D4")
        ws.merge_cells("B3:C4")  # Vedoucí budovy over only part of Alfa, with Pravá ruka under it

    sheet = parse_building_sheet(_sheet(build))
    assert {"building": "Alfa", "room": "A1", "row": "VedouciBudovy"} not in sheet.row_merges
    assert any("nepokrývá celou budovu" in line for line in sheet.warnings)


def test_names_are_matched_to_organizers_by_name_and_the_rest_reported(workspace):
    for name in ("Anna Nováková", "Bob Beran", "Dana Dvorakova", "Hana Horka"):
        mutations.add_organizer(workspace, name)
    state = mutations.get_state(workspace)
    next(o for o in state["organizers"] if o["name"] == "Bob Beran")["cant_attend"] = True
    workspace.save(state)

    _, pending, warnings = mutations.read_building_sheet(workspace, _sheet(_with_leaders))

    by_id = {o["id"]: o["name"] for o in mutations.get_state(workspace)["organizers"]}
    assert [(e["role"], e["building"], e["room"], by_id[e["organizer_id"]]) for e in pending["slots"]] == [
        ("VedouciBudovy", "Alfa", None, "Anna Nováková"),
        ("PravaRuka", "Alfa", None, "Anna Nováková"),
        ("PravaRuka", "Beta", "B1", "Dana Dvorakova"),
        ("TechnickaPodpora", "Alfa", None, "Hana Horka"),
    ]
    assert any("Cyril Cerny" in line and "není mezi organizátory" in line for line in warnings)
    assert any("Bob Beran" in line and "Nemůže se zúčastnit" in line for line in warnings)
    assert pending["row_merges"][0]["row"] == "VedouciBudovy"


def test_saving_a_sheet_replaces_the_leaders_and_merges_in_one_save(workspace, tmp_path):
    ids = {n: mutations.add_organizer(workspace, n)["organizers"][-1]["id"] for n in ("Anna Nováková", "Dana Dvorakova", "Old Boss")}
    mutations.put_config(workspace, [{"name": "Alfa", "capacities": {}, "rooms": [{"name": "A1", "capacities": {}}]}])
    mutations.assign_organizer(workspace, ids["Old Boss"], "VedouciBudovy", "Alfa")
    buildings, pending, _ = mutations.read_building_sheet(workspace, _sheet(_with_leaders))

    with pytest.raises(mutations.ConfirmationRequired) as asked:
        mutations.put_config_from_sheet(workspace, buildings, pending)
    assert asked.value.lines == ["Vedoucí budovy: Old Boss (Alfa)"]
    assert [b["name"] for b in mutations.get_state(workspace)["config"]] == ["Alfa"]  # nothing changed yet

    notes: list[str] = []
    state = mutations.put_config_from_sheet(workspace, buildings, pending, confirmed=True, warnings=notes)

    assert [b["name"] for b in state["config"]] == ["Alfa", "Beta"]
    held = {(e["role"], e["building"], e.get("room"), e["organizer_id"]) for e in state["manual_roles"]["structural"]}
    assert held == {
        ("VedouciBudovy", "Alfa", None, ids["Anna Nováková"]),
        ("PravaRuka", "Alfa", None, ids["Anna Nováková"]),
        ("PravaRuka", "Beta", "B1", ids["Dana Dvorakova"]),
    }
    assert state["row_merges"] == [
        {"building": "Alfa", "room": "A1", "row": "VedouciBudovy"},
        {"building": "Alfa", "room": "A1", "row": "Fotograf"},
    ]
    assert state["cell_merges"]["Fotograf"] == {"Alfa": [["A1", "A2"]], "Beta": [["B1", "B2"]]}
    assert {o["name"]: (o["building"], o["room"]) for o in state["organizers"]} == {
        "Anna Nováková": ("Alfa", None),
        "Dana Dvorakova": ("Beta", "B1"),
        "Old Boss": (None, None),
    }
    assert notes == []


def test_an_organizer_named_in_two_places_is_placed_at_the_last_with_a_note(workspace, tmp_path):
    anna = mutations.add_organizer(workspace, "Anna Nováková")["organizers"][-1]["id"]

    def build(ws) -> None:
        _with_leaders(ws)
        _put(ws, "F9", "Anna Nováková")

    buildings, pending, _ = mutations.read_building_sheet(workspace, _sheet(build))
    notes: list[str] = []
    state = mutations.put_config_from_sheet(workspace, buildings, pending, confirmed=True, warnings=notes)

    assert {(e["role"], e["building"], e.get("room")) for e in state["manual_roles"]["structural"] if e["organizer_id"] == anna} == {
        ("TechnickaPodpora", "Beta", None)
    }
    assert any("Anna Nováková je v tabulce na více místech" in line for line in notes)
    assert state["row_merges"][0]["row"] == "VedouciBudovy"


def test_similar_names_tolerate_nicknames_misspellings_and_partial_names():
    from rostering.organizers import similar_names

    pool = [(1, "Tereza Nováková"), (2, "Jan Svoboda"), (3, "Cyril Černý"), (4, "Petr Novák")]
    assert similar_names("Terka Nováková", pool) == [1]
    assert similar_names("cerny cyril", pool) == [3]  # diacritics, case and word order do not matter
    assert similar_names("Cyrill Cerny", pool) == [3]
    assert similar_names("Svoboda", pool) == [2]  # a surname alone
    assert similar_names("Karel Dvořák", pool) == []


def test_names_without_an_exact_organizer_become_todo_offers_with_similar_candidates(workspace, tmp_path):
    ids = {n: mutations.add_organizer(workspace, n)["organizers"][-1]["id"] for n in ("Anna Nováková", "Cyrill Cerny", "Eva Ehrlich")}
    buildings, pending, warnings = mutations.read_building_sheet(workspace, _sheet(_with_leaders))
    assert {e["name"] for e in pending["unmatched"]} >= {"Cyril Cerny", "Eva Ehrlichova"}
    assert any("Cyril Cerny" in line and "Podobná jména" in line for line in warnings)
    assert any("Bob Beran" in line and "zatím nezařazen" in line for line in warnings)

    mutations.put_config_from_sheet(workspace, buildings, pending, confirmed=True)
    offers = {o["name"]: o for o in mutations.get_organizer_slot_offers(workspace)}

    assert [c["organizer_id"] for c in offers["Cyril Cerny"]["candidates"]] == [ids["Cyrill Cerny"]]
    assert offers["Cyril Cerny"]["label"] == "Pravá ruka (Beta, B1)"
    assert [c["organizer_id"] for c in offers["Eva Ehrlichova"]["candidates"]] == [ids["Eva Ehrlich"]]
    assert offers["Bob Beran"]["candidates"] == []  # nobody resembles them; still listed so they can be added later
    assert "Anna Nováková" not in offers  # matched exactly, so placed rather than offered

    state = mutations.accept_organizer_slot_offer(workspace, offers["Cyril Cerny"]["id"], ids["Cyrill Cerny"])
    assert any(
        (e["role"], e["building"], e.get("room"), e["organizer_id"]) == ("PravaRuka", "Beta", "B1", ids["Cyrill Cerny"])
        for e in state["manual_roles"]["structural"]
    )
    remaining = {o["name"] for o in mutations.get_organizer_slot_offers(workspace)}
    assert "Cyril Cerny" not in remaining

    mutations.dismiss_organizer_slot_offer(workspace, offers["Eva Ehrlichova"]["id"])
    assert "Eva Ehrlichova" not in {o["name"] for o in mutations.get_organizer_slot_offers(workspace)}


def test_an_organizer_added_later_is_offered_for_the_waiting_name(workspace, tmp_path):
    buildings, pending, _ = mutations.read_building_sheet(workspace, _sheet(_with_leaders))
    mutations.put_config_from_sheet(workspace, buildings, pending, confirmed=True)
    assert next(o for o in mutations.get_organizer_slot_offers(workspace) if o["name"] == "Bob Beran")["candidates"] == []

    bob = mutations.add_organizer(workspace, "Bob Beran")["organizers"][-1]["id"]

    offer = next(o for o in mutations.get_organizer_slot_offers(workspace) if o["name"] == "Bob Beran")
    assert [c["organizer_id"] for c in offer["candidates"]] == [bob]
    with pytest.raises(mutations.RosteringError):
        mutations.accept_organizer_slot_offer(workspace, 999, bob)
