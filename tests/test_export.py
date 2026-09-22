import openpyxl

from rostering.domain import (
    Assignment,
    Building,
    Competition,
    Helper,
    ManualRoles,
    OverlayAssignment,
    OverlayRole,
    Role,
    RoleCapacity,
    Room,
    SolveResult,
    StructuralAssignment,
    StructuralRole,
)
from rostering.export.excel import write_roster


def test_write_roster_produces_readable_workbook_with_manual_roles(tmp_path):
    room = Room(name="S3", capacities={Role.Opravovatel: RoleCapacity(1), Role.Zaloha: RoleCapacity(0)})
    building = Building(name="Malá Strana", rooms=[room])
    helpers = [
        Helper(id=1, name="Anna"),
        Helper(id=2, name="Petr"),
    ]
    comp = Competition(buildings={"Malá Strana": building}, helpers=helpers)
    result = SolveResult(
        assignments=[
            Assignment(helper_id=1, helper_name="Anna", building="Malá Strana", room="S3", role=Role.Opravovatel),
            Assignment(helper_id=2, helper_name="Petr", building="Malá Strana", room="S3", role=Role.Zaloha),
        ],
        status="OPTIMAL",
        objective_value=0.0,
    )
    manual = ManualRoles(
        structural=[
            StructuralAssignment(role=StructuralRole.VedouciBudovy, building="Malá Strana", helper_id=1),
        ],
        # Overlay entries have no building of their own — write_roster locates
        # Petr's building via his solved assignment above.
        overlay=[OverlayAssignment(role=OverlayRole.Registrace, helper_id=2)],
    )

    out_path = tmp_path / "roster.xlsx"
    write_roster(comp, result, manual, out_path)

    assert out_path.exists()
    wb = openpyxl.load_workbook(out_path)
    assert wb.sheetnames == ["Pomocníci v místnostech"]

    ws = wb["Pomocníci v místnostech"]
    values = {cell.value for row in ws.iter_rows() for cell in row if cell.value}
    assert "Anna" in values
    assert "Vedoucí budovy" in values
    assert any(cell.value and "Petr" in cell.value for row in ws.iter_rows() for cell in row)


def test_write_roster_merges_cells_within_one_row_only(tmp_path):
    rooms = [
        Room(
            name="R1",
            capacities={Role.Opravovatel: RoleCapacity(1), Role.Menic: RoleCapacity(1)},
        ),
        Room(
            name="R2",
            capacities={Role.Opravovatel: RoleCapacity(1), Role.Menic: RoleCapacity(1)},
        ),
    ]
    building = Building(name="B", rooms=rooms)
    helpers = [
        Helper(id=1, name="Anna"),
        Helper(id=2, name="Petr"),
        Helper(id=3, name="Klara"),
        Helper(id=4, name="David"),
    ]
    comp = Competition(buildings={"B": building}, helpers=helpers)
    result = SolveResult(
        assignments=[
            Assignment(helper_id=1, helper_name="Anna", building="B", room="R1", role=Role.Opravovatel),
            Assignment(helper_id=2, helper_name="Petr", building="B", room="R2", role=Role.Opravovatel),
            # Menic (a different row) is deliberately left unmerged, to
            # confirm the Opravovatel merge doesn't leak into other rows.
            Assignment(helper_id=3, helper_name="Klara", building="B", room="R1", role=Role.Menic),
            Assignment(helper_id=4, helper_name="David", building="B", room="R2", role=Role.Menic),
        ],
        status="OPTIMAL",
        objective_value=0.0,
    )

    out_path = tmp_path / "roster.xlsx"
    # Merge R1/R2 for the Opravovatel row only — like merging cells in
    # Excel, not the whole column.
    write_roster(comp, result, ManualRoles(), out_path, cell_merges={"Opravovatel": {"B": [["R1", "R2"]]}})

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Pomocníci v místnostech"]
    values = [[cell.value for cell in row] for row in ws.iter_rows()]

    # Header room row is untouched by a row-scoped merge — still two
    # separate room columns.
    assert values[1][1] == "R1"
    assert values[1][2] == "R2"

    # Opravovatel row: R1/R2 merged into one cell for this row only — Anna
    # and Petr both land under column 1 (stacked), column 2 is empty
    # (consumed by the merge).
    opravovatel_row_idx = next(i for i, r in enumerate(values) if r[0] == "Opravovatelé")
    assert values[opravovatel_row_idx][1] == "Anna"
    assert values[opravovatel_row_idx + 1][1] == "Petr"
    assert values[opravovatel_row_idx][2] is None
    assert values[opravovatel_row_idx + 1][2] is None

    # Menic row has no merge applied — stays two separate columns.
    menic_row = next(r for r in values if r[0] == "Měniči")
    assert menic_row[1] == "Klara"
    assert menic_row[2] == "David"


def test_write_roster_ignores_stale_cell_merge_pairs(tmp_path):
    room = Room(name="S3", capacities={Role.Zaloha: RoleCapacity(0)})
    building = Building(name="B", rooms=[room])
    helpers = [Helper(id=1, name="Anna")]
    comp = Competition(buildings={"B": building}, helpers=helpers)
    result = SolveResult(
        assignments=[Assignment(helper_id=1, helper_name="Anna", building="B", room="S3", role=Role.Zaloha)],
        status="OPTIMAL",
        objective_value=0.0,
    )

    out_path = tmp_path / "roster.xlsx"
    # "S3"/"S4" aren't adjacent rooms in this building at all — must not raise.
    write_roster(comp, result, ManualRoles(), out_path, cell_merges={"Zaloha": {"B": [["S3", "S4"]]}})
    assert out_path.exists()


def test_write_roster_with_no_manual_roles(tmp_path):
    room = Room(name="S3", capacities={Role.Zaloha: RoleCapacity(0)})
    building = Building(name="B", rooms=[room])
    helpers = [Helper(id=1, name="Anna")]
    comp = Competition(buildings={"B": building}, helpers=helpers)
    result = SolveResult(
        assignments=[Assignment(helper_id=1, helper_name="Anna", building="B", room="S3", role=Role.Zaloha)],
        status="OPTIMAL",
        objective_value=0.0,
    )

    out_path = tmp_path / "roster.xlsx"
    write_roster(comp, result, ManualRoles(), out_path)

    assert out_path.exists()
    openpyxl.load_workbook(out_path)  # does not raise
