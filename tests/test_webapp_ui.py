"""The NiceGUI web app, driven in-process with NiceGUI's user simulation against
a temp-dir Season seeded with synthetic Helpers and Organizers (never data/):
the page, the People tab and person sheet, Tags, Buildings, Solver, the Roster
tab's grid events, the Solve modal and the sidebar."""
import asyncio
import functools
import io

import pytest
from nicegui import ui
from nicegui.testing import User
from nicegui.testing.user_simulation import user_simulation

from rostering.persistence import config_store
from rostering.persistence.workspace import Workspace
from rostering.solver.model import MAX_ROLE_COST, SolverConfig
from rostering.persistence.serialize import solver_config_to_dict
from rostering.webapp import labels, mutations
from rostering.webapp.ui import solving
from rostering.webapp.ui.app import root
from rostering.webapp.ui.tabs.buildings import layout_key

CONFIG = [
    {
        "name": "Karlín",
        "rooms": [
            {"name": "K1", "capacities": {"Zaloha": {"minimum": 0}}},
            {"name": "K2", "capacities": {}},
        ],
        "capacities": {},
    },
    {"name": "Impakt", "rooms": [{"name": "I1", "capacities": {"Zaloha": {"minimum": 0}}}], "capacities": {}},
]


def _helper(helper_id, name, **extra):
    return {
        "id": helper_id,
        "name": name,
        "role_preferences": {},
        "building_preferences": [],
        "friends": [],
        "can_bring_notebook": False,
        "can_bring_camera": False,
        "unresolved_friend_names": [],
        **extra,
    }


@pytest.fixture
def seasons(tmp_path, monkeypatch):
    """A Season with Anna and Bára (Helpers), Boss (Organizer, holding Vedoucí
    budovy at Karlín) and one Tag, "GCHD"."""
    monkeypatch.setenv("ROSTERING_SEASONS_DIR", str(tmp_path / "seasons"))
    # The default layout is a user-level file: keep it out of the repository.
    save_default = config_store.save_default_config
    monkeypatch.setattr(
        config_store, "save_default_config", lambda buildings, path=None: save_default(buildings, path=tmp_path / "b.yaml")
    )
    workspace = Workspace()
    state = workspace.load()
    state["helpers"] = [_helper(1, "Anna"), _helper(2, "Bára")]
    state["next_helper_id"] = 3
    workspace.save(state)
    workspace.create_season("2026-jaro")
    mutations.put_config(workspace, CONFIG)
    mutations.add_organizer(workspace, "Boss")
    mutations.assign_organizer(workspace, 1, "VedouciBudovy", "Karlín")
    mutations.add_tag(workspace, "GCHD")
    return Workspace()


@pytest.fixture
async def user(seasons):
    async with user_simulation(root=root) as user:
        await user.open("/")
        yield user


def _state(workspace) -> dict:
    return mutations.get_state(workspace)


async def _go(user: User, tab: str) -> None:
    """Switch to a step tab and wait for it to be drawn."""
    user.find(marker="step-tabs").elements.pop().value = tab
    await asyncio.sleep(0.2)


def _rows(user: User, marker: str) -> list[dict]:
    return user.find(marker=marker).elements.pop().rows


# ---------------------------------------------------------------------- page
async def test_the_page_shows_the_open_season_and_the_six_steps(user: User):
    await user.should_see("Ročník 2026-jaro")
    assert sorted(t._props["name"] for t in user.find(kind=ui.tab).elements) == sorted(labels.TABS)


async def test_switching_tabs_shows_that_step(user: User):
    await _go(user, labels.TAB_SOLVER)
    await user.should_see("Ceny rolí")


@pytest.fixture
def review_list_reads(monkeypatch) -> list[int]:
    """Counts the reads of the review list, which goes through every stored
    Season (one entry per read)."""
    reads: list[int] = []
    real = mutations.get_uncertain_matches

    @functools.wraps(real)
    def counted(workspace):
        reads.append(1)
        return real(workspace)

    monkeypatch.setattr(mutations, "get_uncertain_matches", counted)
    return reads


async def test_switching_tabs_reads_no_stored_season(user: User, review_list_reads):
    # The Season's data is as it was: the to-do count and panel are left alone.
    for tab in (labels.TAB_TAGS, labels.TAB_ROSTER, labels.TAB_PEOPLE):
        await _go(user, tab)
    assert review_list_reads == []


async def test_a_change_reads_the_review_list_once_for_the_header_and_the_open_panel(
    user: User, seasons, review_list_reads
):
    user.find(marker="todo-button").click()
    await asyncio.sleep(0.2)
    review_list_reads.clear()

    user.find(marker="helper-table").trigger("cant_attend", {"id": 1, "value": True})
    await asyncio.sleep(0.2)
    assert _state(seasons)["helpers"][0]["cant_attend"] is True
    assert review_list_reads == [1]


# ---------------------------------------------------------------------- People
async def test_the_people_tab_lists_organizers_and_helpers(user: User):
    assert [r["name"] for r in _rows(user, "organizer-table")] == ["Boss"]
    assert [r["placement"] for r in _rows(user, "organizer-table")] == ["Karlín"]
    assert [r["name"] for r in _rows(user, "helper-table")] == ["Anna", "Bára"]


async def test_both_tables_load_their_sheet_with_a_header_button(user: User):
    user.find(kind=ui.button, content="Načíst pomocníky")
    user.find(kind=ui.button, content="Načíst organizátory")
    await user.should_not_see("Soubor s odpověďmi")


async def test_with_no_season_open_both_sheets_can_be_loaded_to_create_one(seasons):
    mutations.new_season(seasons)
    async with user_simulation(root=root) as user:
        await user.open("/")
        await user.should_see("Zatím není otevřený žádný ročník.")
        user.find(kind=ui.button, content="Načíst pomocníky")
        user.find(kind=ui.button, content="Načíst organizátory")


async def test_ticking_a_helpers_cant_attend_flags_them(user: User, seasons):
    user.find(marker="helper-table").trigger("cant_attend", {"id": 1, "value": True})
    await asyncio.sleep(0.1)
    assert _state(seasons)["helpers"][0]["cant_attend"] is True
    assert _rows(user, "helper-table")[0]["cant_attend"] is True

    user.find(marker="helper-table").trigger("cant_attend", {"id": 1, "value": False})
    await asyncio.sleep(0.1)
    assert not _state(seasons)["helpers"][0].get("cant_attend")


async def test_flagging_a_placed_organizer_asks_first(user: User, seasons):
    user.find(marker="organizer-table").trigger("cant_attend", {"id": 1, "value": True})
    await user.should_see("Označit jako Nemůže se zúčastnit?")
    assert not _state(seasons)["organizers"][0].get("cant_attend")

    user.find(kind=ui.button, content="Označit jako Nemůže se zúčastnit").click()
    await asyncio.sleep(0.1)
    assert _state(seasons)["organizers"][0]["cant_attend"] is True
    assert _state(seasons)["manual_roles"]["structural"] == []


async def test_a_row_opens_the_person_sheet_where_edits_save(user: User, seasons):
    user.find(marker="helper-table").trigger("rowClick", [{}, {"id": 1}, 0])
    await user.should_see(marker="person-name")
    user.find(marker="person-name").clear().type("Anna K.")
    user.find(marker="person-save").click()
    await asyncio.sleep(0.1)
    assert _state(seasons)["helpers"][0]["name"] == "Anna K."


async def test_tagging_an_organizer_in_their_sheet_saves_at_once(user: User, seasons):
    user.find(marker="organizer-table").trigger("rowClick", [{}, {"id": 1}, 0])
    await user.should_see(marker="person-tags")
    user.find(marker="person-tags").elements.pop().value = [1]
    await asyncio.sleep(0.1)
    assert _state(seasons)["organizers"][0]["tags"] == [1]


async def test_deleting_a_helper_asks_first(user: User, seasons):
    user.find(marker="helper-table").trigger("rowClick", [{}, {"id": 2}, 0])
    await user.should_see(marker="person-delete")
    user.find(marker="person-delete").click()
    await user.should_see("Smazat tohoto pomocníka?")
    assert len(_state(seasons)["helpers"]) == 2

    user.find(marker="confirm-ok").click()
    await asyncio.sleep(0.1)
    assert [h["name"] for h in _state(seasons)["helpers"]] == ["Anna"]


async def test_the_add_organizer_dialog_creates_one(user: User, seasons):
    user.find(marker="add-organizer").click()
    await user.should_see(marker="new-name")
    user.find(marker="new-name").type("Nova")
    user.find(marker="new-submit").click()
    await asyncio.sleep(0.1)
    assert [o["name"] for o in _state(seasons)["organizers"]] == ["Boss", "Nova"]


async def test_the_add_helper_dialog_creates_one(user: User, seasons):
    user.find(marker="add-helper").click()
    await user.should_see(marker="new-contact")
    user.find(marker="new-name").type("Nova Helper")
    user.find(marker="new-contact").type("nova@example.com")
    user.find(marker="new-submit").click()
    await asyncio.sleep(0.1)
    assert [h["name"] for h in _state(seasons)["helpers"]] == ["Anna", "Bára", "Nova Helper"]


async def _friends_tab(user: User, helper_id: int = 1) -> None:
    await user.open("/")
    user.find(marker="helper-table").trigger("open_friends", {"id": helper_id})
    await user.should_see(marker="person-friends")


async def test_unmatched_friend_names_open_the_sheet_on_the_friends_tab(user: User, seasons):
    state = seasons.load()
    state["helpers"][0].update(unresolved_friend_names=["Terka"], friend_name_order=["Terka"])
    seasons.save(state)
    await _friends_tab(user)
    await user.should_see("K přiřazení")
    await user.should_see("“Terka”")
    await user.should_see("Přiřazení kamarádi")


async def test_matching_a_friend_name_saves_the_decision(user: User, seasons):
    state = seasons.load()
    state["helpers"][0].update(unresolved_friend_names=["Bárka"], friend_name_order=["Bárka"])
    seasons.save(state)
    await _friends_tab(user)
    user.find(marker="person-match").elements.pop().value = ["h2"]
    await asyncio.sleep(0.1)
    helper = _state(seasons)["helpers"][0]
    assert helper["friends"] == [2]
    assert helper["unresolved_friend_names"] == []


async def test_the_dismiss_button_disables_the_match_and_toggles_back(user: User, seasons):
    state = seasons.load()
    state["helpers"][0].update(unresolved_friend_names=["Bárka"], friend_name_order=["Bárka"])
    seasons.save(state)
    await _friends_tab(user)
    user.find(marker="person-dismiss").click()
    await asyncio.sleep(0.1)
    helper = _state(seasons)["helpers"][0]
    assert helper["friend_name_decisions"] == {"Bárka": None}
    assert not user.find(marker="person-match").elements.pop().enabled
    user.find(marker="person-dismiss").click()
    await asyncio.sleep(0.1)
    helper = _state(seasons)["helpers"][0]
    assert helper["unresolved_friend_names"] == ["Bárka"]
    assert user.find(marker="person-match").elements.pop().enabled


async def test_the_friends_picker_saves_at_once_and_feeds_the_forced_picker(user: User, seasons):
    await _friends_tab(user)
    user.find(marker="person-friends").elements.pop().value = ["h2"]
    await asyncio.sleep(0.1)
    assert _state(seasons)["helpers"][0]["friends"] == [2]
    forced = user.find(marker="person-forced").elements.pop()
    assert forced.options == {"h2": "Bára"}  # only the Helper's own friends


async def test_the_forced_picker_forces_and_unforces_a_friend(user: User, seasons):
    mutations.update_helper(seasons, 1, friends=[2])
    await _friends_tab(user)
    user.find(marker="person-forced").elements.pop().value = ["h2"]
    await asyncio.sleep(0.1)
    (group,) = _state(seasons)["forced_groups"]
    assert (group["name"], group["rules"]) == ("Anna + Bára", [{"kind": "share", "axis": "room"}])

    user.find(marker="person-forced").elements.pop().value = []
    await asyncio.sleep(0.1)
    assert _state(seasons)["forced_groups"] == []
    assert _state(seasons)["helpers"][0]["friends"] == [2]  # the wish stays


# ---------------------------------------------------------------------- Forced friends
async def test_the_forced_friends_tab_shows_member_states_and_badges(user: User, seasons):
    from rostering.webapp import forced_groups

    people = {h["id"]: h["person_id"] for h in _state(seasons)["helpers"]}
    forced_groups.add_group(seasons, "Rodina", [people[1], people[2]], ["room", "role"])
    mutations.promote_helper(seasons, 1)
    await user.open("/")
    await _go(user, labels.TAB_FORCED)
    await user.should_see("Rodina")
    await user.should_see("Anna (nezařazený organizátor, neaktivní),")
    await user.should_see("Role se na Anna neuplatní")


async def test_the_forced_friends_card_lists_its_rules_and_warns_about_places_it_lacks(user: User, seasons):
    from rostering.forced_friends import Rule
    from rostering.webapp import forced_groups

    people = {h["id"]: h["person_id"] for h in _state(seasons)["helpers"]}
    forced_groups.add_group(
        seasons,
        "Pevná",
        [people[1]],
        [Rule.be("building", ["Karlín"]), Rule.be("role", ["Fotograf"], must=False)],
    )
    state = _state(seasons)
    state["forced_groups"][0]["rules"].append(Rule.be("building", ["Troja"], must=False).to_dict())
    seasons.save(state)
    await user.open("/")
    await _go(user, labels.TAB_FORCED)
    await user.should_see("Členové musí být v budově Karlín")
    await user.should_see("Členové nesmí mít roli Fotograf")
    await user.should_see("budova Troja v tomto ročníku není")


async def test_a_forced_friends_group_is_made_from_a_list_of_rules_in_the_dialog(user: User, seasons):
    people = {h["id"]: h["person_id"] for h in _state(seasons)["helpers"]}
    await user.open("/")
    await _go(user, labels.TAB_FORCED)
    user.find("Nová skupinka").click()
    await user.should_see(marker="group-name")
    user.find(marker="group-name").type("Pevná")
    user.find(marker="group-people").elements.pop().value = [people[1], people[2]]
    user.find(marker="rule-op-0").elements.pop().value = "be_must"  # the first row: musí být v ...
    await asyncio.sleep(0.1)
    user.find(marker="rule-axis-0").elements.pop().value = "building"
    await asyncio.sleep(0.1)
    user.find(marker="rule-values-0").elements.pop().value = ["Karlín"]
    user.find(marker="rule-add").click()
    await asyncio.sleep(0.1)
    user.find(marker="rule-axis-1").elements.pop().value = "room"  # the second row still shares: now a Room
    user.find(marker="group-save").click()
    await asyncio.sleep(0.2)

    (group,) = _state(seasons)["forced_groups"]
    assert group["name"] == "Pevná"
    assert group["rules"] == [
        {"kind": "be", "must": True, "axis": "building", "values": ["Karlín"]},
        {"kind": "share", "axis": "room"},
    ]
    await user.should_see("Členové musí být v budově Karlín")
    await user.should_see("Členové musí sdílet místnost")


async def test_the_dialog_refuses_contradictory_rules_and_keeps_the_group_unsaved(user: User, seasons):
    people = {h["id"]: h["person_id"] for h in _state(seasons)["helpers"]}
    await user.open("/")
    await _go(user, labels.TAB_FORCED)
    user.find("Nová skupinka").click()
    await user.should_see(marker="group-name")
    user.find(marker="group-name").type("Pevná")
    user.find(marker="group-people").elements.pop().value = [people[1]]
    user.find(marker="rule-op-0").elements.pop().value = "be_must"
    await asyncio.sleep(0.1)
    user.find(marker="rule-values-0").elements.pop().value = ["Karlín"]
    user.find(marker="rule-add").click()
    await asyncio.sleep(0.1)
    user.find(marker="rule-op-1").elements.pop().value = "be_not"
    await asyncio.sleep(0.1)
    user.find(marker="rule-values-1").elements.pop().value = ["Karlín"]
    user.find(marker="group-save").click()
    await asyncio.sleep(0.2)

    assert _state(seasons)["forced_groups"] == []
    await user.should_see("si odporují")


# ---------------------------------------------------------------------- Tags
async def test_the_tags_tab_adds_and_removes_carriers_in_one_save(user: User, seasons):
    mutations.set_organizer_tags(seasons, 1, [1])
    await user.open("/")
    await _go(user, labels.TAB_TAGS)
    user.find(marker="tag-table").trigger("rowClick", [{}, {"id": 1}, 0])
    await user.should_see("Kdo štítek nese (1)")
    assert [r["key"] for r in user.find(marker="tag-carriers").elements.pop().selected] == ["o1"]

    carriers = user.find(marker="tag-carriers")
    carriers.trigger("selection", {"added": True, "rows": [{"key": "h1"}], "keys": ["h1"]})
    carriers.trigger("selection", {"added": False, "rows": [], "keys": ["o1"]})
    user.find(marker="tag-carriers-save").click()
    await asyncio.sleep(0.1)
    state = _state(seasons)
    assert mutations.helper_tags(state, 1)["direct"] == [1]
    assert mutations.organizer_tags(state, 1)["direct"] == []


async def test_the_tag_sheet_is_hidden_until_a_row_is_clicked(user: User, seasons):
    await user.open("/")
    await _go(user, labels.TAB_TAGS)
    await user.should_see(marker="tag-table")
    await user.should_not_see(marker="tag-name")
    assert user.find(marker="tag-table").elements.pop().rows[0]["name"] == "GCHD"

    user.find(marker="tag-table").trigger("rowClick", [{}, {"id": 1}, 0])
    await user.should_see(marker="tag-name")
    assert user.find(marker="tag-name").elements.pop().value == "GCHD"


async def test_the_new_tag_button_opens_the_create_form_and_then_the_new_tag(user: User, seasons):
    await user.open("/")
    await _go(user, labels.TAB_TAGS)
    user.find(marker="add-tag").click()
    await user.should_see(marker="tag-name")
    user.find(marker="tag-name").type("Nový")
    user.find(marker="tag-save").click()
    await asyncio.sleep(0.2)
    assert [t["name"] for t in _state(seasons)["tags"]] == ["GCHD", "Nový"]
    await user.should_see("Kdo štítek nese (0)")


async def test_a_tag_gets_its_rules_from_the_same_rule_rows_a_group_uses(user: User, seasons):
    await user.open("/")
    await _go(user, labels.TAB_TAGS)
    user.find(marker="add-tag").click()
    await user.should_see(marker="tag-name")
    user.find(marker="tag-name").type("Jen K1")
    user.find(marker="rule-add").click()  # a new Tag starts with no rule; a Tag cannot share, so the row is "musí být v"
    await asyncio.sleep(0.1)
    ops = user.find(marker="rule-op-0").elements.pop()
    assert "share" not in ops.options and "be_must" in ops.options
    user.find(marker="rule-axis-0").elements.pop().value = "room"
    await asyncio.sleep(0.1)
    user.find(marker="rule-values-0").elements.pop().value = ["KarlínK1"]
    user.find(marker="tag-save").click()
    await asyncio.sleep(0.2)

    tag = next(t for t in _state(seasons)["tags"] if t["name"] == "Jen K1")
    assert tag["rules"] == [{"kind": "be", "must": True, "axis": "room", "values": [["Karlín", "K1"]]}]


# ---------------------------------------------------------------------- Buildings
async def test_buildings_edits_are_an_unsaved_draft_until_saved(user: User, seasons):
    await _go(user, labels.TAB_BUILDINGS)
    await user.should_not_see(marker="unsaved")
    user.find(marker="add-building").click()
    await user.should_see(marker="unsaved")
    assert len(_state(seasons)["config"]) == 2  # nothing saved yet

    user.find(marker="buildings-save").click()
    await user.should_not_see(marker="unsaved")
    names = [b["name"] for b in _state(seasons)["config"]]
    assert len(names) == 3 and len(set(names)) == 3


async def test_reset_restores_the_bundled_layout_as_an_unsaved_draft(user: User, seasons):
    await _go(user, labels.TAB_BUILDINGS)
    user.find(marker="buildings-reset").click()
    await user.should_see(marker="unsaved")
    assert layout_key(_state(seasons)["config"]) == layout_key(CONFIG)  # only the draft changed


async def test_loading_a_sheet_saves_the_layout_after_confirming(user: User, seasons):
    import openpyxl
    from nicegui.elements.upload_files import SmallFileUpload
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"], ws["A3"] = "Nová budova", "X1", "Opravovatelé"
    ws["B3"].fill = PatternFill("solid", fgColor="FFFFE599")
    content = io.BytesIO()
    wb.save(content)

    await _go(user, labels.TAB_BUILDINGS)
    uploader = user.find(marker="buildings-sheet-upload").elements.pop()
    await uploader.handle_uploads([SmallFileUpload(name="rooms.xlsx", content_type="", _data=content.getvalue())])
    await user.should_see("Nahradit vedoucí a sloučené buňky?")  # a sheet replaces the leaders held now
    assert [b["name"] for b in _state(seasons)["config"]] == ["Karlín", "Impakt"]  # not saved until confirmed
    user.find(marker="confirm-ok").click()
    await asyncio.sleep(0.2)
    await user.should_not_see(marker="unsaved")
    saved = _state(seasons)["config"]
    assert [b["name"] for b in saved] == ["Nová budova"]
    assert saved[0]["rooms"] == [{"name": "X1", "capacities": {"Opravovatel": {"minimum": 1}}}]


async def test_saving_a_sheet_with_leaders_asks_before_replacing_the_slots(user: User, seasons):
    import openpyxl
    from nicegui.elements.upload_files import SmallFileUpload
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"], ws["A3"], ws["A4"] = "Nová budova", "X1", "Vedoucí budovy", "Opravovatelé"
    ws["B3"] = "Boss, Neznámý Člověk"
    ws["B4"].fill = PatternFill("solid", fgColor="FFFFE599")
    content = io.BytesIO()
    wb.save(content)

    await _go(user, labels.TAB_BUILDINGS)
    uploader = user.find(marker="buildings-sheet-upload").elements.pop()
    await uploader.handle_uploads([SmallFileUpload(name="rooms.xlsx", content_type="", _data=content.getvalue())])
    await user.should_see("není mezi organizátory")
    await user.should_see("Nahradit vedoucí a sloučené buňky?")
    assert [(e["role"], e["building"]) for e in _state(seasons)["manual_roles"]["structural"]] == [("VedouciBudovy", "Karlín")]

    user.find(marker="confirm-ok").click()
    await asyncio.sleep(0.2)

    state = _state(seasons)
    assert [b["name"] for b in state["config"]] == ["Nová budova"]
    assert [(e["role"], e["building"], e["organizer_id"]) for e in state["manual_roles"]["structural"]] == [("VedouciBudovy", "Nová budova", 1)]
    await user.should_not_see(marker="unsaved")


# ---------------------------------------------------------------------- Solver
async def test_restore_defaults_puts_the_role_costs_back(user: User, seasons):
    await _go(user, labels.TAB_SOLVER)
    user.find(marker="cost-ne").elements.pop().value = 3
    await user.should_see(marker="unsaved")
    user.find(marker="restore-costs").click()
    await asyncio.sleep(0.1)
    defaults = SolverConfig()
    assert user.find(marker="cost-ne").elements.pop().value == defaults.role_costs.ne
    for field in ("ano", "klidne", "nevadi", "spise_ne", "ne", "zaloha"):
        slider = user.find(marker=f"cost-{field}").elements.pop()
        assert (slider.props["min"], slider.props["max"]) == (0, MAX_ROLE_COST)


async def test_a_saved_cost_above_the_maximum_is_shown_at_the_top(user: User, seasons):
    config = solver_config_to_dict(SolverConfig())
    config["role_costs"]["ne"] = 99
    state = seasons.load()
    state["solver_config"] = config
    seasons.save(state)
    await user.open("/")
    await _go(user, labels.TAB_SOLVER)
    assert user.find(marker="cost-ne").elements.pop().value == MAX_ROLE_COST


async def test_the_solver_tab_cannot_be_left_with_unsaved_changes(user: User, seasons):
    await _go(user, labels.TAB_SOLVER)
    user.find(marker="cost-ne").elements.pop().value = 3
    await user.should_see(marker="unsaved")

    await _go(user, labels.TAB_ROSTER)
    await user.should_see("Máte neuložené změny")
    await user.should_see("Ceny rolí")  # still on the Solver tab
    assert user.find(marker="step-tabs").elements.pop().value == labels.TAB_SOLVER


async def test_reverting_the_solver_draft_frees_the_tab_and_shows_the_saved_values(user: User, seasons):
    await _go(user, labels.TAB_SOLVER)
    saved = user.find(marker="cost-ne").elements.pop().value
    user.find(marker="cost-ne").elements.pop().value = 3
    await user.should_see(marker="unsaved")

    user.find(marker="solver-revert").click()
    await asyncio.sleep(0.2)
    await user.should_not_see(marker="unsaved")
    assert user.find(marker="cost-ne").elements.pop().value == saved

    await _go(user, labels.TAB_ROSTER)
    assert user.find(marker="step-tabs").elements.pop().value == labels.TAB_ROSTER


async def test_saving_the_solver_draft_frees_the_tab(user: User, seasons):
    await _go(user, labels.TAB_SOLVER)
    user.find(marker="cost-ne").elements.pop().value = 3
    await user.should_see(marker="unsaved")
    user.find(marker="solver-save").click()
    await asyncio.sleep(0.2)
    assert _state(seasons)["solver_config"]["role_costs"]["ne"] == 3

    await _go(user, labels.TAB_ROSTER)
    assert user.find(marker="step-tabs").elements.pop().value == labels.TAB_ROSTER


# ---------------------------------------------------------------------- Solve modal
async def test_a_finished_solve_closes_its_modal_by_itself(user: User):
    with user:
        result = await solving.run_in_modal("Sestavuji…", lambda: "done")
    assert result == "done"
    await user.should_not_see("Sestavuji…")


@pytest.mark.parametrize(
    "error, message",
    [
        (mutations.RosteringError("No roster found within 60 seconds"), "No roster found within 60 seconds"),
        (RuntimeError("boom"), "Neočekávaná chyba: boom"),
    ],
)
async def test_a_failed_solve_keeps_the_modal_open_until_closed(user: User, error, message):
    runs = []

    def work():
        runs.append(1)
        raise error

    async def solve():
        with user:  # NiceGUI's element context lives per task
            return await solving.run_in_modal("Sestavuji…", work)

    task = asyncio.create_task(solve())
    await user.should_see(message)
    user.find(kind=ui.button, content="Zavřít").click()
    assert await asyncio.wait_for(task, 5) is None
    assert runs == [1]  # closing does not run it again


async def test_solving_over_unlocked_assignments_asks_first_then_solves(user: User, seasons):
    mutations.move_helper(seasons, 1, "Karlín", "K1", "Zaloha")
    await user.open("/")
    await _go(user, labels.TAB_ROSTER)
    user.find(marker="solve").click()
    await user.should_see("Nahradit neuzamčená přiřazení?")
    user.find(marker="confirm-ok").click()
    for _ in range(100):  # the solver runs in a worker thread
        await asyncio.sleep(0.1)
        if len(_state(seasons)["assignments"]) == 2:
            break
    assert {a["helper_id"] for a in _state(seasons)["assignments"]} == {1, 2}


async def test_save_and_solve_asks_over_the_redrawn_tab_then_solves(user: User, seasons):
    # Saving redraws the tab the button sits in; the confirmation must survive that.
    mutations.move_helper(seasons, 1, "Karlín", "K1", "Zaloha")
    await user.open("/")
    await _go(user, labels.TAB_BUILDINGS)
    user.find(marker="save-and-solve").click()
    await user.should_see("Nahradit neuzamčená přiřazení?")
    user.find(marker="confirm-ok").click()
    for _ in range(100):
        await asyncio.sleep(0.1)
        if len(_state(seasons)["assignments"]) == 2:
            break
    assert {a["helper_id"] for a in _state(seasons)["assignments"]} == {1, 2}
    await user.should_see(marker="roster-grid")  # the Roster tab is shown


# ---------------------------------------------------------------------- Roster grid
async def _roster(user: User, seasons):
    mutations.move_helper(seasons, 1, "Karlín", "K1", "Zaloha")
    mutations.move_helper(seasons, 2, "Karlín", "K2", "Zaloha")
    await user.open("/")
    await _go(user, labels.TAB_ROSTER)
    await user.should_see(marker="roster-grid")
    return user.find(marker="roster-grid")


async def test_a_drop_moves_the_helper(user: User, seasons):
    grid = await _roster(user, seasons)
    grid.trigger("helper_drop", {"helper_id": 1, "building": "Impakt", "room": "I1", "role": "Zaloha"})
    await asyncio.sleep(0.1)
    placed = {a["helper_id"]: (a["building"], a["room"]) for a in _state(seasons)["assignments"]}
    assert placed[1] == ("Impakt", "I1")


async def test_moving_a_helper_out_of_their_additional_role_room_asks_first(user: User, seasons):
    grid = await _roster(user, seasons)
    grid.trigger("manual_set", {"key": "UvadeciUcastniku", "building": "Karlín", "room": "K1", "names": ["anna"]})
    await asyncio.sleep(0.1)
    assert [e["helper_id"] for e in _state(seasons)["manual_roles"]["overlay"]] == [1]  # typed name matched

    user.find(marker="roster-grid").trigger("helper_drop", {"helper_id": 1, "building": "Impakt", "room": "I1", "role": "Zaloha"})
    await user.should_see("Odebrat pomocníka z manuální role?")
    user.find(kind=ui.button, content="Přesunout a odebrat").click()
    await asyncio.sleep(0.1)
    state = _state(seasons)
    assert state["manual_roles"]["overlay"] == []
    assert next(a for a in state["assignments"] if a["helper_id"] == 1)["building"] == "Impakt"


async def test_lock_and_organizer_drop_and_merge_events(user: User, seasons):
    grid = await _roster(user, seasons)
    grid.trigger("lock", {"helper_id": 2, "locked": True})
    await asyncio.sleep(0.1)
    assert next(a for a in _state(seasons)["assignments"] if a["helper_id"] == 2).get("locked") is True

    user.find(marker="roster-grid").trigger(
        "organizer_drop",
        {"organizer_id": 1, "key": "VedouciMistnosti", "building": "Karlín", "room": "K2", "source": None},
    )
    await asyncio.sleep(0.1)
    slots = [(s["role"], s.get("room")) for s in _state(seasons)["manual_roles"]["structural"]]
    assert ("VedouciMistnosti", "K2") in slots

    user.find(marker="roster-grid").trigger(
        "cell_merge", {"key": "Zaloha", "building": "Karlín", "pairs": [["K1", "K2"]], "merged": True}
    )
    await asyncio.sleep(0.1)
    assert _state(seasons)["cell_merges"]["Zaloha"]["Karlín"] == [["K1", "K2"]]


async def test_merging_two_leadership_rows_into_a_tall_cell_and_dropping_into_it(user: User, seasons):
    await _roster(user, seasons)  # the grid is redrawn after every change: find it afresh each time
    # Karlín has K1, K2: Pravá ruka over both, so it can meet Vedoucí budovy (Boss holds that).
    user.find(marker="roster-grid").trigger("cell_merge", {"key": "PravaRuka", "building": "Karlín", "pairs": [["K1", "K2"]], "merged": True})
    await asyncio.sleep(0.1)
    user.find(marker="roster-grid").trigger("row_merge", {"key": "VedouciBudovy", "building": "Karlín", "room": "K1", "merged": True})
    await asyncio.sleep(0.1)
    state = _state(seasons)
    assert state["row_merges"] == [{"building": "Karlín", "room": "K1", "row": "VedouciBudovy"}]
    assert {(e["role"], e["organizer_id"]) for e in state["manual_roles"]["structural"]} == {
        ("VedouciBudovy", 1),
        ("PravaRuka", 1),
    }
    # The sideways merge under a tall cell is refused; the cell stays.
    user.find(marker="roster-grid").trigger("cell_merge", {"key": "PravaRuka", "building": "Karlín", "pairs": [["K1", "K2"]], "merged": False})
    await user.should_see("Nejdřív ji rozdělte")
    assert len(_state(seasons)["row_merges"]) == 1

    user.find(marker="roster-grid").trigger("row_merge", {"key": "VedouciBudovy", "building": "Karlín", "room": "K1", "merged": False})
    await asyncio.sleep(0.1)
    state = _state(seasons)
    assert state["row_merges"] == []
    assert {(e["role"], e["organizer_id"]) for e in state["manual_roles"]["structural"]} == {("VedouciBudovy", 1)}


async def test_a_leadership_slot_takes_a_typed_organizer_name(user: User, seasons):
    grid = await _roster(user, seasons)
    grid.trigger("manual_set", {"key": "TechnickaPodpora", "building": "Impakt", "room": None, "names": ["Nový Org"]})
    await asyncio.sleep(0.1)
    assert "Nový Org" in [o["name"] for o in _state(seasons)["organizers"]]


async def test_the_details_card_opens_and_closes_on_the_same_chip(user: User, seasons):
    grid = await _roster(user, seasons)
    click = {"helper_id": 1, "x": 10, "y": 10, "width": 1000, "height": 800}
    grid.trigger("card", click)
    await user.should_see("Preference rolí")
    user.find(marker="roster-grid").trigger("card", click)
    await user.should_not_see("Preference rolí")


async def test_an_organizers_details_card_opens_closes_and_edits_in_the_person_sheet(user: User, seasons):
    mutations.update_organizer(seasons, 1, phone="777 111 222")
    grid = await _roster(user, seasons)
    click = {"organizer_id": 1, "x": 10, "y": 10, "width": 1000, "height": 800}
    grid.trigger("card", click)
    await user.should_see("Telefon")
    await user.should_see("777 111 222")
    await user.should_see("Vedoucí budovy · Karlín")
    user.find(marker="roster-grid").trigger("card", click)  # the same chip again closes it
    await user.should_not_see("777 111 222")

    user.find(marker="roster-grid").trigger("card", click)
    await user.should_see(marker="card-edit")
    user.find(marker="card-edit").click()
    await user.should_see(marker="person-name")
    await user.should_not_see(marker="card-edit")


async def test_the_tags_legend_lists_tags_on_the_grid_and_a_click_filters(user: User, seasons):
    state = _state(seasons)
    tag_id = next(t["id"] for t in state["tags"] if t["name"] == "GCHD")
    mutations.set_helper_tags(seasons, 1, [tag_id])
    await _roster(user, seasons)
    await user.should_not_see("Štítky v mřížce:")  # the Tags overlay is off to begin with

    user.find(kind=ui.chip, content="Štítky").elements.pop().selected = True
    await user.should_see("Štítky v mřížce:")
    await user.should_see(marker=f"tag-pill-{tag_id}")

    user.find(marker=f"tag-pill-{tag_id}").click()
    await asyncio.sleep(0.2)
    picker = user.find(kind=ui.select, content="Filtrovat podle štítků").elements.pop()
    assert picker.value == [tag_id]
    await user.should_see("✓ GCHD")

    user.find(marker=f"tag-pill-{tag_id}").click()  # a second click takes it out again
    await asyncio.sleep(0.2)
    assert picker.value == []


async def test_the_details_card_shows_tags_and_edits_in_the_person_sheet(user: User, seasons):
    tag_id = _state(seasons)["tags"][0]["id"]
    mutations.add_tag_to_helpers(seasons, tag_id, [1])
    grid = await _roster(user, seasons)
    grid.trigger("card", {"helper_id": 1, "x": 10, "y": 10, "width": 1000, "height": 800})
    await user.should_see("Štítky")
    await user.should_see("GCHD")
    user.find(marker="card-edit").click()
    await user.should_see(marker="person-name")
    await user.should_not_see(marker="card-edit")  # the card hands over to the sheet
    user.find(marker="person-name").clear().type("Anna K.")
    user.find(marker="person-save").click()
    await asyncio.sleep(0.1)
    assert _state(seasons)["helpers"][0]["name"] == "Anna K."


# ---------------------------------------------------------------------- sidebar
async def test_a_version_saves_and_restores_after_confirming(user: User, seasons):
    user.find(marker="version-name").type("před úpravou").trigger("keydown.enter")
    await asyncio.sleep(0.1)
    assert [v["name"] for v in mutations.list_versions(seasons)] == ["před úpravou"]

    mutations.add_tag(seasons, "Nový")
    user.find(marker="restore-version").click()
    await user.should_see("Obnovit verzi")
    user.find(marker="confirm-ok").click()
    await asyncio.sleep(0.1)
    assert [tag["name"] for tag in _state(seasons)["tags"]] == ["GCHD"]


async def test_start_over_asks_first(user: User, seasons):
    user.find(marker="start-over").click()
    await user.should_see("Začít znovu?")
    user.find(marker="confirm-ok").click()
    await asyncio.sleep(0.1)
    assert _state(seasons)["helpers"] == []


# ---------------------------------------------------------------------- Organizers' import
async def _upload_organizers(user: User, content: bytes) -> None:
    from nicegui.elements.upload_files import SmallFileUpload

    uploader = user.find(marker="organizers-upload").elements.pop()
    await uploader.handle_uploads([SmallFileUpload(name="organizers.xlsx", content_type="", _data=content)])
    await asyncio.sleep(0.3)


async def test_loading_the_organizers_sheet_fills_the_table_and_the_todo_summary(user: User, seasons):
    from tests.survey_factory import organizer_survey_bytes

    await _upload_organizers(user, organizer_survey_bytes(4, seed=7))
    rows = _rows(user, "organizer-table")
    assert len(rows) == 5  # Boss and the four from the sheet
    loaded = [r for r in rows if r["name"] != "Boss"]
    assert sorted(r["phone"] == "—" for r in loaded) == [False, False, False, True]  # row 2 left it blank
    assert any(r["tshirt"].startswith(("pánské", "dámské")) for r in loaded)
    user.find(marker="todo-button").click()  # the to-do panel is drawn only while open
    await user.should_see("Co změnilo poslední nahrání organizátorů")
    assert _state(seasons)["organizer_upload_summary"]["new"]

    user.find(kind=ui.button, content="Skrýt").click()
    await asyncio.sleep(0.1)
    assert "organizer_upload_summary" not in _state(seasons)


async def test_the_organizer_sheet_shows_the_answers_and_saves_phone_and_size(user: User, seasons):
    from tests.survey_factory import organizer_survey_bytes

    await _upload_organizers(user, organizer_survey_bytes(2, seed=7))
    organizer = _state(seasons)["organizers"][-1]
    user.find(marker="organizer-table").trigger("rowClick", [{}, {"id": organizer["id"]}, 0])
    await user.should_see(marker="person-name")
    user.find(kind=ui.tab, content="Odpovědi").click()
    await user.should_see("Role: Vedoucí místnosti")
    await user.should_see(organizer["survey"]["role_VedouciMistnosti"])

    user.find(marker="person-save").click()
    await asyncio.sleep(0.1)
    saved = next(o for o in _state(seasons)["organizers"] if o["id"] == organizer["id"])
    assert saved["phone"] == organizer["phone"] and saved.get("tshirt_size") == organizer.get("tshirt_size")


async def test_a_returning_organizer_is_offered_for_review_in_the_todo_panel(user: User, seasons):
    from tests.survey_factory import ORGANIZER_HEADERS, generate_organizer_survey

    mutations.new_season(seasons)
    seasons.create_season("2025-podzim")
    mutations.add_organizer(seasons, "Dana Stará")
    mutations.new_season(seasons)
    seasons.create_season("2026-podzim")
    frame = generate_organizer_survey(1, seed=0)
    frame[ORGANIZER_HEADERS["name"]] = ["Dana Stará"]
    buffer = io.BytesIO()
    frame.to_excel(buffer, index=False)
    mutations.import_organizers(seasons, buffer.getvalue(), "organizers.xlsx")
    await user.open("/")
    await user.should_not_see("Možní vracející se organizátoři (1)")  # drawn only while open

    user.find(marker="todo-button").click()
    await user.should_see("Možní vracející se organizátoři (1)")
    user.find(kind=ui.button, content="Propojit").click()
    await asyncio.sleep(0.2)
    organizer = _state(seasons)["organizers"][0]
    assert organizer["link_confirmed"] is True
    assert not mutations.get_uncertain_organizer_matches(seasons)


async def test_the_person_sheet_lists_a_helpers_raw_survey_responses(user: User, seasons):
    state = _state(seasons)
    state["helpers"][0]["survey_responses"] = [
        {"question": "Tvoje jméno a příjmení", "answer": "Anna"},
        {"question": "Něco navíc, co aplikace nečte", "answer": "Mám rád/a koláče"},
        {"question": "Poznámka", "answer": ""},
    ]
    seasons.save(state)
    await user.open("/")  # the open page still holds the state it loaded
    user.find(marker="helper-table").trigger("rowClick", [{}, {"id": 1}, 0])
    await user.should_see(marker="person-name")
    user.find(kind=ui.tab, content="Odpovědi z dotazníku").click()
    await user.should_see("Něco navíc, co aplikace nečte")
    await user.should_see("Mám rád/a koláče")
    await user.should_see("— bez odpovědi —")
