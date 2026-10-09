"""Locked Assignments: the per-Helper lock flag in the Season's saved state,
carried by the Assignment serialization, kept by a full Solve and by hand
moves. Runs against a temp-dir workspace seeded with synthetic Helpers (never
anything from data/)."""
import importlib

import pytest

from rostering.domain import Assignment, Role
from rostering.persistence.serialize import assignment_from_dict, assignment_to_dict
from rostering.persistence.workspace import Workspace
from rostering.webapp import mutations

TWO_ROOMS = [
    {
        "name": "B",
        "rooms": [
            {"name": "R1", "capacities": {"Zaloha": {"minimum": 0}}},
            {"name": "R2", "capacities": {"Zaloha": {"minimum": 0}}},
        ],
        "capacities": {},
    }
]


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ROSTERING_BUILDINGS_CONFIG_PATH", str(tmp_path / "buildings-config.yaml"))
    from rostering.persistence import config_store as config_store_module

    importlib.reload(config_store_module)
    return Workspace(root=tmp_path / "workspace")


def _helper(helper_id: int, name: str, prefs: dict) -> dict:
    return {
        "id": helper_id,
        "name": name,
        "role_preferences": prefs,
        "building_preferences": [],
        "friends": [],
        "can_bring_notebook": True,
        "can_bring_camera": True,
        "unresolved_friend_names": [],
    }


def _seed(workspace: Workspace, config=TWO_ROOMS) -> None:
    """Anna and Petr both dearly want Opravovatel and refuse every other role."""
    prefs = {role: "Ne" for role in ("Menic", "Skenovac", "Kreslic", "Fotograf", "Zaloha")}
    prefs["Opravovatel"] = "Ano"
    state = workspace.load()
    state["helpers"] = [_helper(1, "Anna", dict(prefs)), _helper(2, "Petr", dict(prefs))]
    workspace.save(state)
    mutations.put_config(workspace, config)


def _assignment(state, helper_id):
    return next(a for a in state["assignments"] if a["helper_id"] == helper_id)


def test_the_lock_flag_round_trips_through_assignment_serialization():
    locked = Assignment(helper_id=1, helper_name="A", building="B", room="R1", role=Role.Kreslic, locked=True)

    assert assignment_from_dict(assignment_to_dict(locked)).locked is True


def test_a_saved_assignment_without_a_lock_flag_is_unlocked():
    data = {"helper_id": 1, "helper_name": "A", "building": "B", "room": "R1", "role": "Kreslic"}

    assert assignment_from_dict(data).locked is False
    assert "locked" not in assignment_to_dict(assignment_from_dict(data))


def test_a_placed_helper_can_be_locked_and_unlocked(workspace):
    _seed(workspace)
    mutations.solve(workspace)

    locked = mutations.set_lock(workspace, 1, True)
    assert _assignment(locked, 1).get("locked") is True
    assert not _assignment(locked, 2).get("locked")
    assert mutations.get_state(workspace) == locked  # persisted

    unlocked = mutations.set_lock(workspace, 1, False)
    assert not _assignment(unlocked, 1).get("locked")


def test_only_placed_helpers_can_be_locked(workspace):
    _seed(workspace)  # not solved yet: nobody is placed

    with pytest.raises(mutations.RosteringError):
        mutations.set_lock(workspace, 1, True)
    with pytest.raises(mutations.RosteringError):
        mutations.set_lock(workspace, 999, True)


def test_a_full_solve_keeps_a_locked_assignment_and_replaces_the_others(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    # Anna is hand-placed against her preferences and locked; Petr is placed
    # against his too but left unlocked.
    mutations.move_helper(workspace, 1, "B", "R2", "Kreslic")
    mutations.move_helper(workspace, 2, "B", "R2", "Kreslic")
    mutations.set_lock(workspace, 1, True)

    state = mutations.solve(workspace)

    anna = _assignment(state, 1)
    assert (anna["building"], anna["room"], anna["role"]) == ("B", "R2", "Kreslic")
    assert anna["locked"] is True
    petr = _assignment(state, 2)
    assert petr["role"] == "Opravovatel"
    assert not petr.get("locked")


def test_a_full_solve_with_no_locks_is_unchanged(workspace):
    _seed(workspace)

    state = mutations.solve(workspace)

    assert {a["role"] for a in state["assignments"]} == {"Opravovatel"}
    assert not any(a.get("locked") for a in state["assignments"])


def test_moving_a_locked_helper_moves_the_lock_with_them(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.set_lock(workspace, 1, True)

    state = mutations.move_helper(workspace, 1, "B", "R2", "Zaloha")

    moved = _assignment(state, 1)
    assert (moved["room"], moved["role"]) == ("R2", "Zaloha")
    assert moved["locked"] is True


def test_moving_an_unlocked_helper_never_locks_it(workspace):
    _seed(workspace)
    mutations.solve(workspace)

    state = mutations.move_helper(workspace, 1, "B", "R2", "Zaloha")

    assert not _assignment(state, 1).get("locked")


def test_a_lock_on_a_room_that_no_longer_exists_is_dropped_at_the_next_solve(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.move_helper(workspace, 1, "B", "R2", "Kreslic")
    mutations.set_lock(workspace, 1, True)
    one_room = [{"name": "B", "rooms": [TWO_ROOMS[0]["rooms"][0]], "capacities": {}}]
    mutations.put_config(workspace, one_room)

    state = mutations.solve(workspace)

    anna = _assignment(state, 1)
    assert anna["room"] == "R1"
    assert anna["role"] == "Opravovatel"
    assert not anna.get("locked")


def test_locks_are_saved_and_restored_with_a_version(workspace):
    _seed(workspace)
    workspace.create_season("2026-jaro")
    mutations.solve(workspace)
    mutations.set_lock(workspace, 1, True)
    saved = mutations.save_version(workspace, "Locked")
    mutations.set_lock(workspace, 1, False)

    restored = mutations.restore_version(workspace, saved["slug"])

    assert _assignment(restored, 1).get("locked") is True


# -- Bottom-bar controls and the lock lifecycle (#45) ------------------------


def _one_room_config():
    return [{"name": "B", "rooms": [TWO_ROOMS[0]["rooms"][0]], "capacities": {}}]


def _sheet_text(xlsx_bytes: bytes) -> list:
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    return [[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets]


def test_lock_all_placed_locks_every_placed_helper_and_clear_all_locks_frees_them(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    assert mutations.locked_count(mutations.get_state(workspace)) == 0

    locked = mutations.lock_all_placed(workspace)

    assert all(a.get("locked") is True for a in locked["assignments"])
    assert mutations.locked_count(locked) == 2
    assert mutations.get_state(workspace) == locked  # persisted

    cleared = mutations.clear_all_locks(workspace)

    assert not any(a.get("locked") for a in cleared["assignments"])
    assert mutations.locked_count(cleared) == 0
    assert len(cleared["assignments"]) == 2  # the Assignments themselves stay


def test_clear_roster_removes_every_assignment_and_resets_the_solver_result(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.set_lock(workspace, 1, True)
    mutations.mark_stale(workspace, "a Helper changed")
    before = mutations.get_state(workspace)
    assert before["diagnostics"]["status"] is not None

    cleared = mutations.clear_roster(workspace)

    assert cleared["assignments"] == []
    assert cleared["diagnostics"] == {
        "status": None,
        "objective_value": None,
        "unsatisfied_friend_pairs": [],
        "satisfied_friend_pairs": [],
    }
    assert cleared["stale_reasons"] == []
    assert mutations.locked_count(cleared) == 0
    assert [h["id"] for h in cleared["helpers"]] == [1, 2]  # the Helpers stay
    assert mutations.get_state(workspace) == cleared  # persisted
    assert mutations.broken_rules(cleared) == []  # no roster, nothing broken


def test_clear_roster_then_solve_builds_a_fresh_roster(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.clear_roster(workspace)

    solved = mutations.solve(workspace)

    assert len(solved["assignments"]) == 2
    assert solved["diagnostics"]["status"] is not None


def test_clear_roster_does_nothing_before_the_first_solve(workspace):
    _seed(workspace)

    assert mutations.clear_roster(workspace)["assignments"] == []


def test_the_locked_count_follows_single_locks(workspace):
    _seed(workspace)
    mutations.solve(workspace)

    assert mutations.locked_count(mutations.set_lock(workspace, 1, True)) == 1
    assert mutations.locked_count(mutations.set_lock(workspace, 2, True)) == 2
    assert mutations.locked_count(mutations.set_lock(workspace, 1, False)) == 1


def test_the_bulk_lock_controls_do_nothing_before_the_first_solve(workspace):
    _seed(workspace)

    assert mutations.lock_all_placed(workspace)["assignments"] == []
    assert mutations.clear_all_locks(workspace)["assignments"] == []


def test_a_solve_asks_only_about_the_unlocked_assignments_it_would_replace(workspace):
    _seed(workspace)
    assert mutations.unlocked_assignments_replaced(mutations.get_state(workspace)) == 0  # nothing placed yet

    solved = mutations.solve(workspace)
    assert mutations.unlocked_assignments_replaced(solved) == 2

    assert mutations.unlocked_assignments_replaced(mutations.set_lock(workspace, 1, True)) == 1
    assert mutations.unlocked_assignments_replaced(mutations.lock_all_placed(workspace)) == 0  # nothing to lose


def test_a_lock_the_solve_would_drop_counts_as_a_replaced_assignment(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.move_helper(workspace, 1, "B", "R2", "Kreslic")
    mutations.set_lock(workspace, 1, True)

    state = mutations.put_config(workspace, _one_room_config())

    assert mutations.unlocked_assignments_replaced(state) == 2  # Petr, and Anna whose lock would be dropped


def test_a_solve_reports_the_locks_it_dropped_and_why(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.move_helper(workspace, 1, "B", "R2", "Kreslic")
    mutations.move_helper(workspace, 2, "B", "R2", "Kreslic")
    mutations.lock_all_placed(workspace)
    mutations.put_config(workspace, _one_room_config())

    state = mutations.solve(workspace)

    assert state["diagnostics"]["dropped_locks"] == ["2 zámky zrušeny: Místnost R2 už neexistuje"]
    assert mutations.locked_count(state) == 0


def test_a_solve_names_a_removed_building_and_each_reason_separately(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.move_helper(workspace, 1, "B", "R2", "Kreslic")
    mutations.move_helper(workspace, 2, "B", "R1", "Kreslic")
    mutations.lock_all_placed(workspace)
    other_building = [{"name": "C", "rooms": [{"name": "R2", "capacities": {"Zaloha": {"minimum": 0}}}], "capacities": {}}]
    mutations.put_config(workspace, other_building)

    state = mutations.solve(workspace)

    assert state["diagnostics"]["dropped_locks"] == ["2 zámky zrušeny: Budova B už neexistuje"]


def test_a_solve_that_drops_no_lock_reports_none(workspace):
    _seed(workspace)
    mutations.solve(workspace)
    mutations.set_lock(workspace, 1, True)

    assert mutations.solve(workspace)["diagnostics"]["dropped_locks"] == []


def test_locking_operations_leave_the_broken_rules_and_the_export_untouched(workspace):
    _seed(workspace)
    before = mutations.solve(workspace)
    export_before = _sheet_text(mutations.export_xlsx_bytes(workspace))
    lines_before = [b.line for b in mutations.broken_rules(before)]

    locked = mutations.lock_all_placed(workspace)
    assert [b.line for b in mutations.broken_rules(locked)] == lines_before
    assert _sheet_text(mutations.export_xlsx_bytes(workspace)) == export_before  # no lock in the Excel

    cleared = mutations.clear_all_locks(workspace)
    assert [b.line for b in mutations.broken_rules(cleared)] == lines_before


def test_an_older_version_without_the_flag_restores_as_unlocked(workspace):
    _seed(workspace)
    workspace.create_season("2026-jaro")
    mutations.solve(workspace)
    saved = mutations.save_version(workspace, "Before locking")  # no lock flag anywhere in it
    mutations.lock_all_placed(workspace)

    restored = mutations.restore_version(workspace, saved["slug"])

    assert mutations.locked_count(restored) == 0
    assert len(restored["assignments"]) == 2


def test_start_over_clears_all_locks(workspace):
    _seed(workspace)
    workspace.create_season("2026-jaro")
    mutations.solve(workspace)
    mutations.lock_all_placed(workspace)

    state = mutations.reset_workspace(workspace)

    assert state["assignments"] == []
    assert mutations.locked_count(state) == 0


def test_locks_stay_in_their_own_season(workspace):
    _seed(workspace)
    workspace.create_season("2026-jaro")
    mutations.solve(workspace)
    mutations.lock_all_placed(workspace)

    fresh = mutations.new_season(workspace)

    assert fresh["assignments"] == []
    assert mutations.locked_count(fresh) == 0
